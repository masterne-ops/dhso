#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""地市月报渲染 tool —— 把【单市数据 + LLM 分析】渲染成 HTML + docx。

复用 render.py 的公共助手(pct/cls/h/tier_label/_blocks/analysis_html/CSS),
只重写模块①(该市趋势 + 区县下钻);模块②③与全省版同构,渲染逻辑一致。

输入 JSON（--in）结构：
{
  "地市": "杭州市", "本月": "2026-05", "趋势月份": [...],
  "模块1_趋势": <fetch city-report 的 模块1_趋势>,
  "模块1_区县": <... 模块1_区县>,
  "模块2_下钻": <... 模块2_下钻>,        # = drilldown_city 输出(代理商/等级/跑动推广会)
  "模块3_专项": <... 模块3_专项>,
  "分析": {                               # ← LLM 填写,字符串数组(一事一段, 1./2.→有序列表)
     "总览": [...], "区县": [...], "下钻": [...], "专项": [...], "下阶段计划": [...]
  }
}

用法：python3 render_city.py --in city.json --outdir 分析报告 [--basename 杭州月报-2026-05]
"""
import argparse
import json
from pathlib import Path

# 复用全省版渲染助手(同目录)
from render import pct, cls, h, tier_label, analysis_html, CSS  # noqa: E402


def _ol_or_p_docx(doc, qn, Pt, RGBColor, value, color=(0x1f, 0x9d, 0x55), prefix='🔍 '):
    """docx 版分析渲染(独立实现,避免依赖 render.py 内部闭包)。"""
    from render import _blocks
    blocks = _blocks(value)
    first = True
    for kind, payload in blocks:
        if kind == 'ol':
            for i, x in enumerate(payload, 1):
                _para(doc, qn, Pt, RGBColor, f'{i}. {x}', color=color, indent=12)
        else:
            _para(doc, qn, Pt, RGBColor, (prefix if (first and prefix) else '') + payload, color=color)
        first = False


def _para(doc, qn, Pt, RGBColor, txt, size=10.5, color=None, bold=False, indent=0):
    p = doc.add_paragraph()
    run = p.add_run(txt)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    try:
        run.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    return p


# ────────────────────────── HTML ──────────────────────────
def data_table(headers, rows, label_first=True):
    """同时输出桌面表格(.tbl) + 手机卡片(.cards)两套 DOM。

    headers: ['列名', ...]；rows: [[(显示值, css类), ...], ...] 或 [[显示值,...],...]
    每个单元格可为 str 或 (str, css_class) 元组。
    label_first=True 时,卡片标题取每行第 0 列,其余列作 key:value。
    """
    def cell_text(c):
        return c[0] if isinstance(c, tuple) else c

    def cell_cls(c):
        return c[1] if isinstance(c, tuple) and len(c) > 1 else ''

    out = []
    # 桌面表格
    out.append('<table class="tbl"><thead><tr>')
    for i, hd in enumerate(headers):
        out.append(f'<th class="l">{h(hd)}</th>' if i == 0 and label_first else f'<th>{h(hd)}</th>')
    out.append('</tr></thead><tbody>')
    for row in rows:
        out.append('<tr>')
        for i, c in enumerate(row):
            is_label = (i == 0 and label_first)
            classes = ('l' if is_label else 'num')
            if cell_cls(c):
                classes += ' ' + cell_cls(c)
            out.append(f'<td class="{classes}">{cell_text(c)}</td>')
        out.append('</tr>')
    out.append('</tbody></table>')
    # 手机卡片
    out.append('<div class="cards">')
    for row in rows:
        out.append('<div class="dcard">')
        out.append(f'<div class="nm">{cell_text(row[0])}</div>' if label_first else '')
        start = 1 if label_first else 0
        for i in range(start, len(row)):
            c = row[i]
            out.append(f'<div class="kv"><span class="k">{h(headers[i])}</span>'
                       f'<span class="v {cell_cls(c)}">{cell_text(c)}</span></div>')
        out.append('</div>')
    out.append('</div>')
    return ''.join(out)


def render_html(J):
    MS = J['趋势月份']
    CITY = J['地市']
    CUR = J['本月']
    t = J['模块1_趋势']
    dist = J['模块1_区县']
    drill = J['模块2_下钻']
    m3 = J['模块3_专项']
    A = J.get('分析', {})

    P = []
    P.append('<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">')
    P.append('<meta name="viewport" content="width=device-width,initial-scale=1.0">')
    P.append(f'<title>{h(CITY)}月度经营报告 · {CUR}</title><style>{CSS}')
    # ── 工具按钮（导出长图 / 切换手机版），截图与打印时隐藏 ──
    P.append('.toolbar{position:fixed;top:14px;right:14px;z-index:999;display:flex;gap:8px}'
             '.toolbar button{background:#fff;color:var(--pri-d);border:1px solid var(--pri);border-radius:8px;'
             'padding:7px 12px;font-size:13px;font-weight:600;cursor:pointer;box-shadow:0 2px 8px rgba(20,40,80,.15)}'
             '.toolbar button:hover{background:var(--pri);color:#fff}'
             '.toolbar button:disabled{opacity:.6;cursor:wait}'
             '.exporting .toolbar{display:none!important}'
             '@media print{.toolbar{display:none!important}}')
    # ── 卡片视图（手机用）：默认隐藏，窄屏或强制手机版时显示 ──
    P.append('.cards{display:none}'
             '.dcard{border:1px solid var(--line);border-radius:9px;padding:9px 11px;margin:7px 0;background:#fafbfd}'
             '.dcard .nm{font-weight:700;font-size:13.5px;margin-bottom:5px}'
             '.dcard .kv{display:flex;justify-content:space-between;font-size:12.5px;padding:2px 0;border-top:1px dashed var(--line)}'
             '.dcard .kv:first-of-type{border-top:none}'
             '.dcard .kv .k{color:var(--mute)} .dcard .kv .v{font-variant-numeric:tabular-nums;font-weight:600}')
    # ── 强制手机版（点按钮）：窄宽 + 表格转卡片 ──
    P.append('body.mobile .wrap{max-width:440px}'
             'body.mobile header .inner{max-width:440px}'
             'body.mobile .tbl{display:none}'
             'body.mobile .cards{display:block}')
    # ── 自适应：真实窄屏(≤520px)也自动切卡片 ──
    P.append('@media(max-width:520px){.wrap{max-width:100%;padding:0 12px 60px}'
             '.tbl{display:none}.cards{display:block}h1{font-size:20px}}')
    P.append('</style></head><body>')
    P.append('<div class="toolbar">'
             '<button id="viewBtn" onclick="toggleMobile()">📱 手机版</button>'
             '<button id="exportBtn" onclick="exportLongImage()">📷 导出长图</button></div>')
    P.append(f'<header><div class="inner"><h1>{h(CITY)}月度经营报告（诊断版）</h1>')
    P.append(f'<div class="sub">本月 {CUR} · 趋势窗口 {MS[0]}~{MS[-1]} · 口径：全量感知 SO（最新分销价/万）</div>')
    P.append('</div></header><div class="wrap">')

    # 总览
    if A.get('总览'):
        ov = analysis_html(A['总览'], css_class='card', prefix='')
        ov = ov.replace('<p>', '<p><b class="hl">核心结论：</b>', 1)
        P.append(ov)

    # ① 该市趋势 + 区县
    P.append('<h2>① SO 趋势与区县下钻</h2><div class="card">')
    P.append(f"<p>{h(CITY)}（万元）：{MS[0]} <b class='num'>{t['分月SO'].get(MS[0])}</b> → "
             f"{MS[1]} <b class='num'>{t['分月SO'].get(MS[1])}</b> → {MS[2]} <b class='num'>{t['分月SO'].get(MS[2])}</b>　"
             f"本月同比 <span class='{cls(t['本月同比'])}'>{pct(t['本月同比'])}</span>　"
             f"环比 <span class='{cls(t['本月环比'])}'>{pct(t['本月环比'])}</span>"
             f"（去年同月 {t['去年同月SO']}）</p>")
    # 区县表
    P.append('<h3>区县 SO 趋势</h3>')
    P.append(data_table(
        ['区县'] + MS + ['本月同比', '本月环比'],
        [[h(r['区县'])] + [r['分月SO'].get(m, 0) for m in MS] +
         [(pct(r['本月同比']), cls(r['本月同比'])), (pct(r['本月环比']), cls(r['本月环比']))]
         for r in dist['区县明细']]))
    gain = '；'.join(f"{x['区县']}（{pct(x['本月环比'])}）" for x in dist['增长_top'])
    drop = '；'.join(f"{x['区县']}（{pct(x['本月环比'])}）" for x in dist['下滑_top'])
    P.append(f"<p>🟢 <b>增长前三</b>：{gain}</p><p>🔴 <b>下滑前三</b>：{drop}</p>")
    if A.get('区县'):
        P.append(analysis_html(A['区县']))
    P.append('</div>')

    # ② 下钻（代理商/等级/跑动）—— 与全省版同构
    P.append(f'<h2>② {h(CITY)}经营下钻</h2><div class="card">')
    if A.get('下钻'):
        P.append(analysis_html(A['下钻']))
    da = drill['代理商']
    P.append(f"<h3>a. 问题代理商 <span class='mut'>（出货口径，门槛≥{da.get('_体量门槛_万')}万，"
             f"{da.get('_代理商总数')}家中{da.get('_达门槛数')}家达标）</span></h3>")
    P.append(data_table(
        ['代理商'] + [f'{m}出货万' for m in MS] + ['出货环比', '本月交易服务商', '服务商数环比'],
        [[h(x['代理商'])] + [x['出货万_分月'].get(m, 0) for m in MS] +
         [(pct(x['出货环比']), cls(x['出货环比'])), x['本月交易服务商数'],
          (f"{x['服务商数环比']:+d}", cls(x['服务商数环比']))]
         for x in da['问题代理商']]))
    P.append('<h3>b. 服务商等级结构 <span class="mut">（货值口径·红包SO万）</span></h3>')
    struct = drill['等级结构']['分月等级结构']
    P.append(data_table(
        ['等级'] + MS + ['环比', '本月服务商数'],
        [[tier_label(row['等级'])] + [struct[m][row['等级']]['SO万'] for m in MS] +
         [(pct(row['环比']), cls(row['环比'])), row['本月服务商数']]
         for row in drill['等级结构']['等级环比']]))
    for lv, reps in drill['等级结构']['代表服务商'].items():
        if reps:
            names = '；'.join(f"{h(r['客户名称'])}（{h(r['区县'])}，掉{r['掉幅万']}万）" for r in reps[:5])
            P.append(f"<p class='mut'><b>{tier_label(lv)} 档下滑代表：</b>{names}</p>")
    vp = drill['跑动推广会']
    P.append('<h3>c. 跑动 &amp; 推广会</h3>')
    P.append(data_table(
        ['指标'] + MS,
        [['有效打卡'] + [vp['跑动_分月'][m]['有效打卡'] for m in MS],
         ['活跃业务员'] + [vp['跑动_分月'][m]['活跃业务员'] for m in MS],
         ['推广会参会人次'] + [vp['推广会_分月'][m]['参会人次'] for m in MS]]))
    vc = vp['跑动_有效打卡环比']
    if vp.get('推广会_本月有数据'):
        pn = f"推广会参会环比 <span class='{cls(vp['推广会_参会环比'])}'>{pct(vp['推广会_参会环比'])}</span>"
    else:
        pn = f"<span class='a'>推广会本月无数据（截至 {vp.get('推广会_数据末月')}）</span>"
    P.append(f"<p class='mut'>跑动有效打卡环比 <span class='{cls(vc)}'>{pct(vc)}</span>；{pn}</p></div>")

    # ③ 专项在该市
    P.append(f'<h2>③ {h(CITY)}三大产品专项</h2><div class="card">')
    P.append(data_table(
        ['专项'] + [f'{m}出货万' for m in MS] + ['本月环比', '状态'],
        [[x['专项']] + [x['分月出货万'].get(m, 0) for m in MS] +
         [(pct(x['本月环比']), cls(x['本月环比'])),
          ('<span class="pill p-red">下降</span>' if x['是否下降'] else '<span class="pill p-green">增长</span>')]
         for x in m3['专项']]))
    if A.get('专项'):
        P.append(analysis_html(A['专项']))
    for x in m3['专项']:
        if x.get('下钻'):
            dd = x['下钻']
            P.append(f"<h3>🔴 {x['专项']} 下降下钻</h3>")
            P.append('<p class="mut"><b>下降代理商：</b>' + '；'.join(
                f"{h(r['代理商'])}({r['环比_万']:+}万)" for r in dd['代理商下降'][:5] if r['环比_万'] < 0) + '</p>')
            P.append('<p class="mut"><b>下降服务商：</b>' + '；'.join(
                f"{h(r['客户名称'])}({h(r['区县'])},掉{r['掉幅万']}万)" for r in dd['服务商下降'][:6]) + '</p>')
    P.append('</div>')

    # ④ 下阶段计划
    plans = A.get('下阶段计划', [])
    if plans:
        P.append('<h2>④ 下阶段计划</h2>')
        P.append(analysis_html(plans, css_class='card', prefix=''))

    P.append(f'<footer>{h(CITY)}月度经营报告（诊断版）· 数据源：生产库 product_flow_v / install_redpack_v / '
             'visit_record_v / promotion_meeting / product_focus · 数字由 report_tools 计算，分析由 AI 撰写</footer>')
    P.append('</div>')
    # ── 切换手机版 + 导出长图 ──
    fbase = json.dumps(f"{CITY}月度经营报告-{CUR}", ensure_ascii=False)
    P.append('<script src="https://cdn.jsdelivr.net/npm/html2canvas@1.4.1/dist/html2canvas.min.js"></script>')
    P.append('<script>')
    # 切换手机版/桌面版
    P.append('function toggleMobile(){'
             'var m=document.body.classList.toggle("mobile");'
             'document.getElementById("viewBtn").textContent=m?"🖥 桌面版":"📱 手机版";'
             '}')
    # 导出长图：截 header+wrap，宽度跟随当前视图（手机版=窄长图）
    P.append('async function exportLongImage(){')
    P.append('  var btn=document.getElementById("exportBtn");')
    P.append('  if(typeof html2canvas==="undefined"){alert("长图组件未加载，请联网后刷新重试");return;}')
    P.append('  var mobile=document.body.classList.contains("mobile");')
    P.append('  btn.disabled=true;btn.textContent="生成中…";document.body.classList.add("exporting");')
    P.append('  await new Promise(function(r){setTimeout(r,60);});')  # 等隐藏 toolbar 生效
    P.append('  try{')
    P.append('    var canvas=await html2canvas(document.body,{scale:2,useCORS:true,backgroundColor:"#f4f6fa",'
             'width:document.body.clientWidth,windowWidth:document.body.clientWidth});')
    P.append('    var a=document.createElement("a");')
    P.append('    a.download=' + fbase + '+(mobile?"-手机版":"")+".png";')
    P.append('    a.href=canvas.toDataURL("image/png");a.click();')
    P.append('  }catch(e){alert("导出失败："+e.message);}')
    P.append('  finally{document.body.classList.remove("exporting");btn.disabled=false;btn.textContent="📷 导出长图";}')
    P.append('}')
    P.append('</script>')
    P.append('</body></html>')
    return '\n'.join(P)


# ────────────────────────── docx ──────────────────────────
def render_docx(J, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    MS = J['趋势月份']
    CITY = J['地市']
    CUR = J['本月']
    t = J['模块1_趋势']
    dist = J['模块1_区县']
    drill = J['模块2_下钻']
    m3 = J['模块3_专项']
    A = J.get('分析', {})

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(10.5)
    try:
        style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass

    def heading(txt, level=1):
        p = doc.add_paragraph()
        run = p.add_run(txt)
        run.bold = True
        run.font.size = Pt(15 - level * 1.5)
        run.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
        try:
            run.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        except Exception:
            pass

    def para(txt, **kw):
        return _para(doc, qn, Pt, RGBColor, txt, **kw)

    def analysis(value, color=(0x1f, 0x9d, 0x55), prefix='🔍 '):
        _ol_or_p_docx(doc, qn, Pt, RGBColor, value, color=color, prefix=prefix)

    def table(headers, rows):
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = 'Light Grid Accent 1'
        for i, hd in enumerate(headers):
            c = tb.rows[0].cells[i]
            c.text = str(hd)
            for r in c.paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(9)
        for row in rows:
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = '' if v is None else str(v)
                for r in cells[i].paragraphs[0].runs:
                    r.font.size = Pt(9)

    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(f'{CITY}月度经营报告（诊断版）· {CUR}')
    tr.bold = True
    tr.font.size = Pt(18)
    tr.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
    try:
        tr.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass
    para(f'趋势窗口 {MS[0]}~{MS[-1]} · 口径：全量感知 SO（万）', size=9, color=(0x6b, 0x76, 0x86))

    if A.get('总览'):
        heading('核心结论', 2)
        analysis(A['总览'], color=None, prefix='')

    # ①
    heading('① SO 趋势与区县下钻', 1)
    para(f"{CITY}：{MS[0]} {t['分月SO'].get(MS[0])} → {MS[1]} {t['分月SO'].get(MS[1])} → "
         f"{MS[2]} {t['分月SO'].get(MS[2])} 万　本月同比 {pct(t['本月同比'])}　环比 {pct(t['本月环比'])}", bold=True)
    table(['区县'] + MS + ['本月同比', '本月环比'],
          [[r['区县']] + [r['分月SO'].get(m, 0) for m in MS] + [pct(r['本月同比']), pct(r['本月环比'])]
           for r in dist['区县明细']])
    para('🟢 增长前三：' + '；'.join(f"{x['区县']}（{pct(x['本月环比'])}）" for x in dist['增长_top']))
    para('🔴 下滑前三：' + '；'.join(f"{x['区县']}（{pct(x['本月环比'])}）" for x in dist['下滑_top']))
    if A.get('区县'):
        analysis(A['区县'])

    # ②
    heading(f'② {CITY}经营下钻', 1)
    if A.get('下钻'):
        analysis(A['下钻'])
    da = drill['代理商']
    para(f"a. 问题代理商（门槛≥{da.get('_体量门槛_万')}万，{da.get('_代理商总数')}家中{da.get('_达门槛数')}家达标）", bold=True)
    table(['代理商'] + [f'{m}出货万' for m in MS] + ['出货环比', '本月交易服务商', '服务商数环比'],
          [[x['代理商']] + [x['出货万_分月'].get(m, 0) for m in MS] +
           [pct(x['出货环比']), x['本月交易服务商数'], f"{x['服务商数环比']:+d}"] for x in da['问题代理商']])
    para('b. 服务商等级结构（货值口径·红包SO万）', bold=True)
    struct = drill['等级结构']['分月等级结构']
    table(['等级'] + MS + ['环比', '本月服务商数'],
          [[tier_label(row['等级'])] + [struct[m][row['等级']]['SO万'] for m in MS] +
           [pct(row['环比']), row['本月服务商数']] for row in drill['等级结构']['等级环比']])
    for lv, reps in drill['等级结构']['代表服务商'].items():
        if reps:
            para(f"  {tier_label(lv)} 档下滑代表：" + '；'.join(
                f"{r['客户名称']}（{r['区县']}，掉{r['掉幅万']}万）" for r in reps[:5]),
                size=9, color=(0x6b, 0x76, 0x86))
    vp = drill['跑动推广会']
    para('c. 跑动 & 推广会', bold=True)
    table(['指标'] + MS,
          [['有效打卡'] + [vp['跑动_分月'][m]['有效打卡'] for m in MS],
           ['活跃业务员'] + [vp['跑动_分月'][m]['活跃业务员'] for m in MS],
           ['推广会参会人次'] + [vp['推广会_分月'][m]['参会人次'] for m in MS]])
    vc = vp['跑动_有效打卡环比']
    note = (f"推广会参会环比 {pct(vp['推广会_参会环比'])}" if vp.get('推广会_本月有数据')
            else f"推广会本月无数据（截至 {vp.get('推广会_数据末月')}）")
    para(f"跑动有效打卡环比 {pct(vc)}；{note}", size=9, color=(0x6b, 0x76, 0x86))

    # ③
    heading(f'③ {CITY}三大产品专项', 1)
    table(['专项'] + [f'{m}出货万' for m in MS] + ['本月环比', '状态'],
          [[x['专项']] + [x['分月出货万'].get(m, 0) for m in MS] +
           [pct(x['本月环比']), '下降' if x['是否下降'] else '增长'] for x in m3['专项']])
    if A.get('专项'):
        analysis(A['专项'])
    for x in m3['专项']:
        if x.get('下钻'):
            dd = x['下钻']
            heading(f"🔴 {x['专项']} 下降下钻", 2)
            para('下降代理商：' + '；'.join(
                f"{r['代理商']}({r['环比_万']:+}万)" for r in dd['代理商下降'][:5] if r['环比_万'] < 0), size=9)
            para('下降服务商：' + '；'.join(
                f"{r['客户名称']}({r['区县']},掉{r['掉幅万']}万)" for r in dd['服务商下降'][:6]), size=9)

    # ④
    plans = A.get('下阶段计划', [])
    if plans:
        heading('④ 下阶段计划', 1)
        analysis(plans, color=None, prefix='')

    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--in', dest='inp', required=True)
    ap.add_argument('--outdir', required=True)
    ap.add_argument('--basename', default=None)
    ap.add_argument('--no-docx', action='store_true')
    args = ap.parse_args()

    J = json.load(open(args.inp, encoding='utf-8'))
    base = args.basename or f"{J['地市']}月度经营报告-{J['本月']}"
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    (outdir / f"{base}.html").write_text(render_html(J), encoding='utf-8')
    print(f"✓ HTML: {outdir / (base + '.html')}")
    if not args.no_docx:
        render_docx(J, str(outdir / f"{base}.docx"))
        print(f"✓ docx: {outdir / (base + '.docx')}")


if __name__ == '__main__':
    main()
