#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""地市周报渲染 → HTML(手机版+导出长图) + docx。复用 render.py / render_city.py 助手。"""
import argparse
import json
from pathlib import Path

from render import pct, cls, h, CSS, analysis_html  # noqa: E402
from render_city import data_table  # 桌面表格+手机卡片双输出  # noqa: E402


def rate(x):
    return f'{x*100:.0f}%' if x is not None else 'N/A'


# ────────── HTML ──────────
def render_html(J):
    CITY = J['地市']; t = J['①周趋势']; dist = J['②区县']
    drill = J['③代理商跑动']; f = J['④专项']; pm = J['推广会']; A = J.get('分析', {}); pk = J.get('服务商KPI', {})
    actv2 = J.get('激活V2', {}); vsplit = J.get('跑动分组', {}); pbydealer = J.get('服务商按代理商', {})
    L = t['标签']
    P = []
    P.append('<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">')
    P.append('<meta name="viewport" content="width=device-width,initial-scale=1.0">')
    P.append(f'<title>{h(CITY)}经营周报 · {J["本周区间"]}</title><style>{CSS}')
    P.append('.toolbar{position:fixed;top:14px;right:14px;z-index:999;display:flex;gap:8px}'
             '.toolbar button{background:#fff;color:var(--pri-d);border:1px solid var(--pri);border-radius:8px;'
             'padding:7px 12px;font-size:13px;font-weight:600;cursor:pointer;box-shadow:0 2px 8px rgba(20,40,80,.15)}'
             '.toolbar button:hover{background:var(--pri);color:#fff}'
             '.toolbar button:disabled{opacity:.6;cursor:wait}'
             '.exporting .toolbar{display:none!important}@media print{.toolbar{display:none!important}}')
    P.append('.cards{display:none}'
             '.dcard{border:1px solid var(--line);border-radius:9px;padding:9px 11px;margin:7px 0;background:#fafbfd}'
             '.dcard .nm{font-weight:700;font-size:13.5px;margin-bottom:5px}'
             '.dcard .kv{display:flex;justify-content:space-between;font-size:12.5px;padding:2px 0;border-top:1px dashed var(--line)}'
             '.dcard .kv:first-of-type{border-top:none}.dcard .kv .k{color:var(--mute)}'
             '.dcard .kv .v{font-variant-numeric:tabular-nums;font-weight:600}')
    P.append('body.mobile .wrap{max-width:440px}body.mobile header .inner{max-width:440px}'
             'body.mobile .tbl{display:none}body.mobile .cards{display:block}')
    P.append('@media(max-width:520px){.wrap{max-width:100%;padding:0 12px 60px}.tbl{display:none}.cards{display:block}h1{font-size:20px}}')
    # 进度条样式
    P.append('.bar{background:#eef1f6;border-radius:5px;height:16px;position:relative;overflow:hidden;margin:3px 0}'
             '.bar>i{display:block;height:100%;background:linear-gradient(90deg,#2f80ed,#1f5fbf)}'
             '.bar>.t{position:absolute;top:-1px;height:18px;border-left:2px dashed #d64545}')
    P.append('</style></head><body>')
    P.append('<div class="toolbar"><button id="viewBtn" onclick="toggleMobile()">📱 手机版</button>'
             '<button id="exportBtn" onclick="exportLongImage()">📷 导出长图</button></div>')
    P.append(f'<header><div class="inner"><h1>{h(CITY)}经营周报</h1>')
    P.append(f'<div class="sub">本周 {J["本周区间"]}（7天）· 口径：全量感知 SO（万）；专项=上线台数</div>')
    P.append('</div></header><div class="wrap">')

    if A.get('总览'):
        ov = analysis_html(A['总览'], css_class='card', prefix='')
        ov = ov.replace('<p>', '<p><b class="hl">核心结论：</b>', 1)
        P.append(ov)

    # ① 周趋势 + 完成率 + YTD
    P.append('<h2>① 本市 SO 周趋势</h2><div class="card">')
    P.append(f"<p>本周 SO <b class='num'>{t['本周SO']}</b> 万，环比上周 "
             f"<span class='{cls(t['本周环比'])}'>{pct(t['本周环比'])}</span>；全省第 <b>{t['全省位次']}</b>/{t['全省地市数']} 位。</p>")
    P.append(data_table(['周窗口'] + L, [['SO万'] + t['分周SO']], label_first=True))
    # 当月完成率进度条
    rt = t['当月完成率'] or 0
    tp = t['当月时间进度'] or 0
    P.append(f"<h3>当月完成率（{t['当月']}）</h3>")
    P.append(f"<div class='bar'><i style='width:{min(rt*100,100):.0f}%'></i>"
             f"<span class='t' style='left:{min(tp*100,100):.0f}%'></span></div>")
    P.append(f"<p class='mut'>当月累计 {t['当月累计']} 万 / 应达成 {t['当月应达成']} 万，"
             f"完成率 <b>{rate(t['当月完成率'])}</b>；时间进度 {rate(t['当月时间进度'])}（红虚线）——"
             f"完成率达到时间进度即为跟上节奏。</p>")
    # YTD
    P.append(f"<p>📊 <b>当月同比</b>：本月至今（{t['当月今年区间']}）{t['当月今年']} 万，"
             f"较去年同期 {t['当月去年']} 万 "
             f"<span class='{cls(t['当月同比'])}'>{pct(t['当月同比'])}</span>。</p>")
    P.append(f"<p>📈 <b>YTD 累计同比</b>：今年至今 {t['YTD今年']} 万，较去年同期 {t['YTD去年']} 万 "
             f"<span class='{cls(t['YTD累计同比'])}'>{pct(t['YTD累计同比'])}</span>。</p>")
    if A.get('趋势'):
        P.append(analysis_html(A['趋势']))
    P.append('</div>')

    # ② 区县
    P.append('<h2>② 区县周动态</h2><div class="card">')
    P.append(data_table(['区县', '本周SO万', '上周SO万', '本周环比'],
             [[h(r['区县']), r['本周SO'], r['上周SO'], (pct(r['本周环比']), cls(r['本周环比']))]
              for r in dist['区县明细']]))
    P.append('<p>🟢 增长前三：' + '；'.join(f"{x['区县']}（{pct(x['本周环比'])}）" for x in dist['增长_top']) + '</p>')
    P.append('<p>🔴 下滑前三：' + '；'.join(f"{x['区县']}（{pct(x['本周环比'])}）" for x in dist['下滑_top']) + '</p>')
    if A.get('区县'):
        P.append(analysis_html(A['区县']))
    P.append('</div>')

    # ③ 代理商 + 跑动
    P.append('<h2>③ 代理商周动态 &amp; 跑动</h2><div class="card">')
    ds = drill.get('代理商SO', [])
    if ds:
        P.append('<h3>各代理商本周出货（按本周 SO 降序，环比涨绿跌红）</h3>')
        P.append(data_table(['代理商', '本周万', '上周万', '环比'],
                 [[h(x['代理商']), x['本周'], x['上周'], (pct(x['环比']), cls(x['环比']))] for x in ds]))
    else:
        P.append('<p class="mut">本周无代理商出货达统计门槛（≥3 万）。</p>')
    v = drill['跑动']
    P.append(f"<p class='mut'>周跑动合计：打卡 {v['上周打卡']}→{v['本周打卡']}"
             f"（<span class='{cls(v['环比'])}'>{pct(v['环比'])}</span>），本周活跃业务员 {v['本周活跃业务员']} 人。</p>")
    # 大华 vs 代理商业务员分组
    if vsplit:
        dh, ag = vsplit.get('大华', {}), vsplit.get('代理商', {})
        P.append('<h3>跑动人员分类（大华 / 代理商业务员）</h3>')
        P.append(data_table(['人员', '本周打卡', '本周人数', '上周打卡', '打卡环比'],
                 [['🏢 大华业务员', dh.get('本周打卡', 0), dh.get('本周人数', 0), dh.get('上周打卡', 0),
                   (pct(dh.get('打卡环比')), cls(dh.get('打卡环比')))],
                  ['🏪 代理商业务员', ag.get('本周打卡', 0), ag.get('本周人数', 0), ag.get('上周打卡', 0),
                   (pct(ag.get('打卡环比')), cls(ag.get('打卡环比')))]]))
    if A.get('代理商'):
        P.append(analysis_html(A['代理商']))
    P.append('</div>')

    # ④ 服务商发展
    if pk:
        P.append('<h2>④ 服务商发展（本周签约/激活 + 任务达成）</h2><div class="card">')
        actcur = actv2.get('本周达V2', pk['本周激活'])
        P.append(f"<p>本周新签约 <b>{pk['本周签约']}</b> 家；本周激活（本年累计破 V2）<b>{actcur}</b> 家"
                 f"，环比 <span class='{cls(actv2.get('环比'))}'>{pct(actv2.get('环比'))}</span>。</p>")
        if actv2:
            P.append('<h3>激活（达 V2）近 6 周趋势</h3>')
            P.append(data_table(['周窗口'] + actv2['标签'], [['达V2家数'] + actv2['分周达V2']], label_first=True))
        td = pk['任务达成']
        P.append('<h3>任务达成（当前实际 / 年度目标）</h3>')
        P.append(data_table(['指标', '当前实际', '年度目标', '达成率'],
                 [[k, td[k]['实际'], td[k]['目标'], (rate(td[k]['达成率']),
                   cls((td[k]['达成率'] or 0) - 0.999))] for k in ['签约数', '新签', 'V2', 'V3', 'V4']]))
        # 按代理商对比(该市每个代理商的拓客/育客)
        if pbydealer.get('代理商'):
            P.append('<h3>各代理商服务商发展（本周签约 / 新达V2）</h3>')
            P.append(data_table(['代理商', '本周签约', '本周激活(达V2)'],
                     [[h(x['代理商']), x['本周签约'], x['本周激活']]
                      for x in pbydealer['代理商']]))
            P.append("<p class='mut'>代理商口径=上级分销商/所属一级客户。</p>")
        P.append("<p class='mut'>本周激活＝服务商管理表激活时间落本周(是否激活=Y)；V2/V3/V4 为官方等级家数 vs 年度目标。</p>")
        if A.get('服务商'):
            P.append(analysis_html(A['服务商']))
        P.append('</div>')

    # ⑤ 专项 + 推广会
    P.append('<h2>⑤ 三大产品专项周趋势（台数）</h2><div class="card">')
    P.append(data_table(['专项'] + L + ['本周环比'],
             [[x['专项']] + x['分周台数'] + [(pct(x['本周环比']), cls(x['本周环比']))] for x in f['专项']]))
    b, p = pm['本周'], pm['上周']
    P.append('<h3>推广会（本周）</h3>')
    P.append(data_table(['指标', '本周', '上周'],
             [['场次', b['场次'], p['场次']], ['参会人次', b['参会人次'], p['参会人次']],
              ['参会服务商', b['参会服务商'], p['参会服务商']]], label_first=True))
    if A.get('专项'):
        P.append(analysis_html(A['专项']))
    P.append('</div>')

    # ⑥ NP 转入客户进展
    npt = J.get('NP转入')
    if npt:
        P.append('<h2>⑥ NP 转入客户进展</h2><div class="card">')
        _n = npt['转入'] or 0

        def _npr(v):
            return f"{v / _n * 100:.1f}%" if _n else '—'
        P.append(data_table(['漏斗阶段', '家数', '占转入'],
                 [['转入', npt['转入'], '100%' if _n else '—'],
                  ['已报备', npt['已报备'], _npr(npt['已报备'])],
                  ['与我司相关', npt['相关'], _npr(npt['相关'])],
                  ['已签约', npt['已签约'], _npr(npt['已签约'])],
                  ['已激活', npt['已激活'], _npr(npt['已激活'])]], label_first=True))
        P.append(f"<p class='mut'>NP 转入客户转化漏斗，截至最新快照（{npt.get('数据时点', '')}）当前状态；"
                 "已签约=命中签约表，已激活=累计上线≥1000 元。旧批「NP 流转名单」无报备/相关字段，报备/相关仅反映新批。</p>")
        P.append('</div>')

    if A.get('下阶段计划'):
        P.append('<h2>⑦ 下阶段计划</h2>')
        P.append(analysis_html(A['下阶段计划'], css_class='card', prefix=''))

    P.append(f'<footer>{h(CITY)}经营周报 · 本周 {J["本周区间"]} · 数据源：生产库（product_flow_v 等）· '
             '数字由 report_tools 计算</footer>')
    P.append('</div>')
    fbase = json.dumps(f"{CITY}经营周报-{J['本周区间'].replace(' ~ ','至').replace('-','')}", ensure_ascii=False)
    P.append('<script src="https://cdn.jsdelivr.net/npm/html2canvas@1.4.1/dist/html2canvas.min.js"></script>')
    P.append('<script>')
    P.append('function toggleMobile(){var m=document.body.classList.toggle("mobile");'
             'document.getElementById("viewBtn").textContent=m?"🖥 桌面版":"📱 手机版";}')
    P.append('async function exportLongImage(){var btn=document.getElementById("exportBtn");'
             'if(typeof html2canvas==="undefined"){alert("长图组件未加载，请联网刷新");return;}'
             'var mobile=document.body.classList.contains("mobile");'
             'btn.disabled=true;btn.textContent="生成中…";document.body.classList.add("exporting");'
             'await new Promise(function(r){setTimeout(r,60);});'
             'try{var c=await html2canvas(document.body,{scale:2,useCORS:true,backgroundColor:"#f4f6fa",'
             'width:document.body.clientWidth,windowWidth:document.body.clientWidth});'
             'var a=document.createElement("a");a.download=' + fbase + '+(mobile?"-手机版":"")+".png";'
             'a.href=c.toDataURL("image/png");a.click();}catch(e){alert("导出失败："+e.message);}'
             'finally{document.body.classList.remove("exporting");btn.disabled=false;btn.textContent="📷 导出长图";}}')
    P.append('</script></body></html>')
    return '\n'.join(P)


# ────────── docx ──────────
def render_docx(J, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    CITY = J['地市']; t = J['①周趋势']; dist = J['②区县']
    drill = J['③代理商跑动']; f = J['④专项']; pm = J['推广会']; A = J.get('分析', {}); pk = J.get('服务商KPI', {})
    actv2 = J.get('激活V2', {}); vsplit = J.get('跑动分组', {}); pbydealer = J.get('服务商按代理商', {})
    L = t['标签']
    doc = Document()
    st = doc.styles['Normal']; st.font.name = '宋体'; st.font.size = Pt(10.5)
    try:
        st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass

    def para(txt, size=10.5, color=None, bold=False):
        p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
        if color:
            r.font.color.rgb = RGBColor(*color)
        try:
            r.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        except Exception:
            pass
        return p

    def heading(txt, lv=1):
        p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True
        r.font.size = Pt(15 - lv * 1.5); r.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
        try:
            r.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        except Exception:
            pass

    def analysis(value):
        from render import _blocks
        for kind, payload in _blocks(value):
            if kind == 'ol':
                for i, x in enumerate(payload, 1):
                    para(f'{i}. {x}', color=(0x1f, 0x9d, 0x55))
            else:
                para('🔍 ' + payload, color=(0x1f, 0x9d, 0x55))

    def table(headers, rows):
        tb = doc.add_table(rows=1, cols=len(headers)); tb.style = 'Light Grid Accent 1'
        for i, hd in enumerate(headers):
            cell = tb.rows[0].cells[i]; cell.text = str(hd)
            for rr in cell.paragraphs[0].runs:
                rr.bold = True; rr.font.size = Pt(9)
        for row in rows:
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = '' if v is None else str(v)
                for rr in cells[i].paragraphs[0].runs:
                    rr.font.size = Pt(9)

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(f"{CITY}经营周报 · {J['本周区间']}")
    tr.bold = True; tr.font.size = Pt(18); tr.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
    try:
        tr.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass
    para(f"本周7天 · 口径：{J['口径说明']}", size=9, color=(0x6b, 0x76, 0x86))

    heading('① 本市 SO 周趋势', 1)
    para(f"本周 SO {t['本周SO']} 万，环比上周 {pct(t['本周环比'])}；全省第 {t['全省位次']}/{t['全省地市数']} 位。", bold=True)
    table(['周窗口'] + L, [['SO万'] + t['分周SO']])
    para(f"当月完成率（{t['当月']}）：累计 {t['当月累计']} 万 / 应达成 {t['当月应达成']} 万 = "
         f"{rate(t['当月完成率'])}；时间进度 {rate(t['当月时间进度'])}（完成率达时间进度即跟上节奏）。")
    para(f"当月同比：本月至今（{t['当月今年区间']}）{t['当月今年']} 万，较去年同期 {t['当月去年']} 万 {pct(t['当月同比'])}。", bold=True)
    para(f"YTD 累计同比：今年至今 {t['YTD今年']} 万，较去年同期 {t['YTD去年']} 万 {pct(t['YTD累计同比'])}。", bold=True)
    if A.get('趋势'):
        analysis(A['趋势'])

    heading('② 区县周动态', 1)
    table(['区县', '本周SO万', '上周SO万', '本周环比'],
          [[r['区县'], r['本周SO'], r['上周SO'], pct(r['本周环比'])] for r in dist['区县明细']])
    para('🟢 增长前三：' + '；'.join(f"{x['区县']}（{pct(x['本周环比'])}）" for x in dist['增长_top']))
    para('🔴 下滑前三：' + '；'.join(f"{x['区县']}（{pct(x['本周环比'])}）" for x in dist['下滑_top']))
    if A.get('区县'):
        analysis(A['区县'])

    heading('③ 代理商周动态 & 跑动', 1)
    ds = drill.get('代理商SO', [])
    if ds:
        para('各代理商本周出货（按本周 SO 降序）：')
        table(['代理商', '本周万', '上周万', '环比'],
              [[x['代理商'], x['本周'], x['上周'], pct(x['环比'])] for x in ds])
    else:
        para('本周无代理商出货达统计门槛（≥3 万）。')
    v = drill['跑动']
    para(f"周跑动合计：打卡 {v['上周打卡']}→{v['本周打卡']}（{pct(v['环比'])}），本周活跃业务员 {v['本周活跃业务员']} 人。", size=9)
    if vsplit:
        dh, ag = vsplit.get('大华', {}), vsplit.get('代理商', {})
        para('跑动人员分类（大华 / 代理商业务员）：', bold=True)
        table(['人员', '本周打卡', '本周人数', '上周打卡', '打卡环比'],
              [['大华业务员', dh.get('本周打卡', 0), dh.get('本周人数', 0), dh.get('上周打卡', 0), pct(dh.get('打卡环比'))],
               ['代理商业务员', ag.get('本周打卡', 0), ag.get('本周人数', 0), ag.get('上周打卡', 0), pct(ag.get('打卡环比'))]])
    if A.get('代理商'):
        analysis(A['代理商'])

    if pk:
        heading('④ 服务商发展（本周签约/激活 + 任务达成）', 1)
        actcur = actv2.get('本周达V2', pk['本周激活'])
        para(f"本周新签约 {pk['本周签约']} 家；本周激活（本年累计破 V2）{actcur} 家，环比 {pct(actv2.get('环比'))}。", bold=True)
        if actv2:
            para('激活（达 V2）近 6 周趋势：')
            table(['周窗口'] + actv2['标签'], [['达V2家数'] + actv2['分周达V2']])
        td = pk['任务达成']
        para('任务达成（当前实际 / 年度目标）：')
        table(['指标', '当前实际', '年度目标', '达成率'],
              [[k, td[k]['实际'], td[k]['目标'], rate(td[k]['达成率'])] for k in ['签约数', '新签', 'V2', 'V3', 'V4']])
        if pbydealer.get('代理商'):
            para('各代理商服务商发展（本周签约 / 新达V2）：', bold=True)
            table(['代理商', '本周签约', '本周激活(达V2)'],
                  [[x['代理商'], x['本周签约'], x['本周激活']]
                   for x in pbydealer['代理商']])
        para('本周激活＝服务商管理表激活时间落本周(是否激活=Y)；V2/V3/V4 为官方等级家数 vs 年度目标。', size=9, color=(0x6b, 0x76, 0x86))
        if A.get('服务商'):
            analysis(A['服务商'])

    heading('⑤ 三大产品专项周趋势（台数）', 1)
    table(['专项'] + L + ['本周环比'],
          [[x['专项']] + x['分周台数'] + [pct(x['本周环比'])] for x in f['专项']])
    b, p = pm['本周'], pm['上周']
    para('推广会（本周/上周）：', bold=True)
    table(['指标', '本周', '上周'],
          [['场次', b['场次'], p['场次']], ['参会人次', b['参会人次'], p['参会人次']],
           ['参会服务商', b['参会服务商'], p['参会服务商']]])
    if A.get('专项'):
        analysis(A['专项'])

    # ⑥ NP 转入客户进展
    npt = J.get('NP转入')
    if npt:
        heading('⑥ NP 转入客户进展', 1)
        _n = npt['转入'] or 0

        def _npr(v):
            return f"{v / _n * 100:.1f}%" if _n else '—'
        table(['漏斗阶段', '家数', '占转入'],
              [['转入', npt['转入'], '100%' if _n else '—'],
               ['已报备', npt['已报备'], _npr(npt['已报备'])],
               ['与我司相关', npt['相关'], _npr(npt['相关'])],
               ['已签约', npt['已签约'], _npr(npt['已签约'])],
               ['已激活', npt['已激活'], _npr(npt['已激活'])]])
        para('NP 转入客户转化漏斗，截至最新快照当前状态；已签约=命中签约表，已激活=累计上线≥1000元。旧批无报备/相关字段。',
             size=9, color=(0x6b, 0x76, 0x86))

    if A.get('下阶段计划'):
        heading('⑦ 下阶段计划', 1)
        analysis(A['下阶段计划'])

    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--in', dest='inp', required=True)
    ap.add_argument('--outdir', required=True)
    ap.add_argument('--no-docx', action='store_true')
    a = ap.parse_args()
    J = json.load(open(a.inp, encoding='utf-8'))
    base = f"{J['地市']}经营周报-{J['本周区间'].replace(' ~ ', '至').replace('-', '')}"
    outdir = Path(a.outdir); outdir.mkdir(parents=True, exist_ok=True)
    (outdir / f"{base}.html").write_text(render_html(J), encoding='utf-8')
    print(f"✓ HTML: {outdir / (base + '.html')}")
    if not a.no_docx:
        render_docx(J, str(outdir / f"{base}.docx"))
        print(f"✓ docx: {outdir / (base + '.docx')}")


if __name__ == '__main__':
    main()
