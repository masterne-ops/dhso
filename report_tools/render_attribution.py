#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""产品视角经营归因报告渲染 — attribution.json(含 AI 归因) → HTML 手机版 + docx。

复用 render.py 的 CSS/格式化/analysis_html + render_city.py 的 data_table/docx helper。

报告结构:① 总览归因(综述+换代主线)→ ② 经营背景速览 → ③ 🟢 亮点切片 → ④ 🔴 问题切片。
每个切片卡片 = 增量大字 + 定性标签(良性🔄/恶性⚠️) + 归因归咎 + 关注点 + [折叠]下钻&9因素佐证。

用法:
  python3 render_attribution.py --in attribution.json --outdir 产品归因 \\
      [--basename 产品归因-全省-2026-01_2026-04] [--no-docx]
输出: {outdir}/{basename}.html + {basename}.docx
"""
import argparse
import json
from pathlib import Path

from render import CSS, pct, cls, h, analysis_html  # noqa: E402
from render_city import data_table, _para, _ol_or_p_docx  # noqa: E402


# ────────────────────────── 数字格式化 ──────────────────────────
def _f(x, fmt):
    if x is None:
        return '·'
    try:
        return format(float(x), fmt)
    except (ValueError, TypeError):
        return str(x)


def wan(x):
    return _f(x, ',.1f')          # 1.5 万


def swan(x):
    return _f(x, '+,.1f')         # +1.5 / -1.5 万


def i(x):
    if x is None:
        return '·'
    try:
        return f'{int(round(float(x))):,}'
    except (ValueError, TypeError):
        return str(x)


def yoy(x):
    return f"<span class='{cls(x)}'>{pct(x)}</span>"


_BENIGN = ['良性', '换代', '顺势', '收尾', '正常', '季节', '主动收缩',
           '健康', '规模化', '放量', '增长', '储备', '铺垫', '爬坡', '培育', '承接']
_MALIGN = ['恶性', '执行', '掉队', '萎缩', '介入', '流失', '放弃', '退出', '失守', '失控']


def severity(tag, 定性):
    """返回 (pill_class, label)。亮点恒绿;问题按定性文本判良性/恶性/中性。
    良性优先(含良性词即归良性);非空但未命中任何关键词 → 中性「关注」,不武断打红。"""
    t = 定性 or ''
    if tag == '亮点子系列':
        return ('p-green', '🟢 亮点')
    if any(k in t for k in _BENIGN):
        return ('p-amber', '🔄 良性/换代')
    if any(k in t for k in _MALIGN):
        return ('p-red', '⚠️ 需介入')
    if t:
        return ('p-amber', '🔄 关注')
    return ('p-red', '🔴 问题')


# 库龄长 → 标红提示
def _stale(days):
    try:
        return float(days) >= 365
    except (ValueError, TypeError):
        return False


# ────────────────────────── 局部 CSS（叠加在 render.CSS 上）──────────────────────────
EXTRA_CSS = (
    '.toolbar{position:fixed;top:14px;right:14px;z-index:999;display:flex;gap:8px}'
    '.toolbar button{background:#fff;color:var(--pri-d);border:1px solid var(--pri);border-radius:8px;'
    'padding:7px 12px;font-size:13px;font-weight:600;cursor:pointer;box-shadow:0 2px 8px rgba(20,40,80,.15)}'
    '.toolbar button:hover{background:var(--pri);color:#fff}'
    '.toolbar button:disabled{opacity:.6;cursor:wait}'
    '.exporting .toolbar{display:none!important}'
    '@media print{.toolbar{display:none!important}}'
    '.cards{display:none}'
    '.dcard{border:1px solid var(--line);border-radius:9px;padding:9px 11px;margin:7px 0;background:#fafbfd}'
    '.dcard .nm{font-weight:700;font-size:13.5px;margin-bottom:5px}'
    '.dcard .kv{display:flex;justify-content:space-between;font-size:12.5px;padding:2px 0;border-top:1px dashed var(--line)}'
    '.dcard .kv:first-of-type{border-top:none}'
    '.dcard .kv .k{color:var(--mute)} .dcard .kv .v{font-variant-numeric:tabular-nums;font-weight:600}'
    'body.mobile .wrap{max-width:440px}body.mobile header .inner{max-width:440px}'
    'body.mobile .tbl{display:none}body.mobile .cards{display:block}'
    '@media(max-width:520px){.wrap{max-width:100%;padding:0 12px 60px}'
    '.tbl{display:none}.cards{display:block}h1{font-size:20px}}'
    # 切片卡片
    '.slice .shead{display:flex;align-items:center;gap:9px;flex-wrap:wrap;margin-bottom:4px}'
    '.slice .sname{font-size:16px;font-weight:700}'
    '.slice .metric{margin-left:auto;font-size:20px;font-weight:800;font-variant-numeric:tabular-nums}'
    '.slice .metric.g{color:var(--green)}.slice .metric.r{color:var(--red)}'
    '.slice details{margin-top:8px;border-top:1px dashed var(--line);padding-top:6px}'
    '.slice summary{cursor:pointer;color:var(--pri-d);font-size:12.5px;font-weight:600}'
    '.dim{font-size:13px;margin:12px 0 5px;color:var(--pri-d)}'
    '.stale{color:var(--red);font-weight:700}'
)


# ────────────────────────── HTML ──────────────────────────
def _background_html(bg):
    if not bg:
        return ''
    out = ['<div class="card">']
    pm = bg.get('推广会', {}) or {}
    if pm.get('当期'):
        c, y = pm['当期'], pm.get('同期', {})
        out.append(f"<p>🎤 <b>推广会</b>:当期 {i(c.get('场次'))} 场 / 签到 {i(c.get('签到公司'))} 家"
                   f"(同期 {i(y.get('场次'))} 场),会后红包 {i(c.get('会后红包金额'))} 元 / "
                   f"夜视王上线 {i(c.get('夜视王上线台数'))} 台。</p>")
    vs = bg.get('跑动', {}) or {}
    if vs.get('当期'):
        c, y = vs['当期'], vs.get('同期', {})
        out.append(f"<p>👣 <b>跑动</b>:当期 {i(c.get('拜访次数'))} 次 / 覆盖 {i(c.get('覆盖服务商'))} 家服务商"
                   f"(同期 {i(y.get('拜访次数'))} 次),大华自打卡 {i(c.get('大华打卡次数'))} 次。</p>")
    sr = bg.get('业务员红包', {}) or {}
    if sr.get('当期'):
        c = sr['当期']
        out.append(f"<p>🧧 <b>业务员红包</b>:使用 {i(c.get('红包使用金额'))} 元 / 发放 {i(c.get('发放客户数'))} 客户,"
                   f"转化激活 {i(c.get('转化激活客户'))} 户,平均解锁率 {pct(c.get('平均解锁率'))},"
                   f"带来上线 {i(c.get('带来上线金额'))} 元。</p>")
    sf = bg.get('门头', {}) or {}
    if sf.get('投入服务商') is not None:
        out.append(f"<p>🏠 <b>门头</b>({h(sf.get('年份'))}年,年度滞后):投入 {i(sf.get('投入服务商'))} 家 / "
                   f"{wan(sf.get('投入金额_万'))} 万,激活 {i(sf.get('激活数'))} 家,"
                   f"带来上线 {wan(sf.get('带来上线金额_万'))} 万。</p>")
    out.append('</div>')
    return ''.join(out)


def _drill_rows(d, dim, n=3):
    """合并跌幅Top + 涨幅Top,去重,按增量升序。返回 data_table 行。"""
    seen = {}
    for r in (d.get('跌幅Top', [])[:n] + d.get('涨幅Top', [])[:n]):
        name = r.get('服务商') if dim == '服务商' else r.get('维度')
        if name and name not in seen:
            seen[name] = r
    rows = []
    for r in sorted(seen.values(), key=lambda r: (r.get('增量_万') or 0)):
        name = r.get('服务商') if dim == '服务商' else r.get('维度')
        if dim == '服务商' and r.get('代理商'):
            name = f"{name}（{r['代理商']}）"
        inc = r.get('增量_万')
        rows.append([h(str(name)), (swan(inc), cls(inc)), wan(r.get('货值_万'))])
    return rows


def _drill_html(drill):
    if not drill:
        return ''
    out = []
    for dim in ['地市', '代理商', '服务商']:
        d = drill.get(dim, {}) or {}
        rows = _drill_rows(d, dim)
        if rows:
            out.append(f'<div class="dim">{dim} · 涨跌贡献</div>')
            out.append(data_table([dim, '增量万', '货值万'], rows))
    return ''.join(out)


def _factor_html(f):
    if not f:
        return ''
    out = ['<div class="dim">9 因素</div>']
    rows = []
    pu = f.get('铺货', {})
    if isinstance(pu, dict) and pu.get('当期'):
        c, y = pu['当期'], pu.get('同期', {})
        rows.append(['铺货台数', i(c.get('铺货台数')), i(y.get('铺货台数'))])
        rows.append(['铺货金额万', wan(c.get('铺货金额_万')), wan(y.get('铺货金额_万'))])
        rows.append(['覆盖下级服务商', i(c.get('覆盖下级服务商')), i(y.get('覆盖下级服务商'))])
    rp = f.get('安装红包', {})
    if isinstance(rp, dict) and rp.get('当期'):
        c, y = rp['当期'], rp.get('同期', {})
        rows.append(['红包上线台数', i(c.get('上线台数')), i(y.get('上线台数'))])
        rows.append(['红包中奖金额', i(c.get('中奖金额')), i(y.get('中奖金额'))])
        rows.append(['红包覆盖服务商', i(c.get('覆盖服务商')), i(y.get('覆盖服务商'))])
    if rows:
        out.append(data_table(['因素指标', '当期', '同期'], rows))
    iv = f.get('备货', {})
    if isinstance(iv, dict) and iv.get('在库台数') is not None:
        days = iv.get('平均库龄天')
        dtxt = f"<span class='stale'>{i(days)} 天</span>" if _stale(days) else f"{i(days)} 天"
        out.append(f"<p class='mut'>📦 备货({h(iv.get('盘库季度'))}):在库 {i(iv.get('在库台数'))} 台 / "
                   f"{wan(iv.get('在库货值_万'))} 万 · 平均库龄 {dtxt}</p>")
    pp = f.get('服务商画像', {})
    if isinstance(pp, dict) and pp.get('购买服务商总数') is not None:
        out.append(f"<p class='mut'>👤 画像:购买 {i(pp.get('购买服务商总数'))} 家 · 竞品TOP {i(pp.get('竞品TOP数'))} · "
                   f"新签 {i(pp.get('新签数'))} · 激活 {i(pp.get('激活数'))} · "
                   f"交易频次 {pp.get('平均本年交易频次')} · 未交易 {i(pp.get('平均未交易天数'))} 天</p>")
        cd = pp.get('客户分类分布')
        if isinstance(cd, dict) and cd:
            top = sorted(cd.items(), key=lambda kv: -(kv[1] or 0))[:4]
            out.append('<p class="mut">　客户结构:' + ' · '.join(f"{h(k)} {i(v)}" for k, v in top) + '</p>')
    return ''.join(out)


def _slice_card(s, tag):
    g = s.get('归因') or {}
    sev_cls, sev_label = severity(tag, g.get('定性', ''))
    inc = s.get('增量_万')
    mc = 'g' if (inc or 0) >= 0 else 'r'
    out = ['<div class="card slice">']
    out.append(f'<div class="shead"><span class="sname">{h(s.get("子系列") or "(未命名子系列)")}</span>'
               f'<span class="pill {sev_cls}">{sev_label}</span>'
               f'<span class="metric {mc}">{swan(inc)}万</span></div>')
    out.append(f"<p class='mut'>货值 {wan(s.get('货值_万'))} 万 · 同比 {yoy(s.get('同比'))} · "
               f"环比 {yoy(s.get('环比'))} · {i(s.get('台数'))} 台</p>")
    if g.get('定性'):
        out.append(f'<p class="lead"><b class="hl">判定:</b>{h(g["定性"])}</p>')
    if g.get('归因归咎'):
        out.append('<h3>归因 / 归咎</h3>')
        out.append(analysis_html(g['归因归咎'], css_class='analysis', prefix=''))
    if g.get('关注点'):
        out.append('<h3>关注点 / 行动</h3>')
        out.append(analysis_html(g['关注点'], css_class='callout', prefix='👉 '))
    drill_h = _drill_html(s.get('下钻', {}))
    factor_h = _factor_html(s.get('因素', {}))
    if drill_h or factor_h:
        out.append('<details><summary>📊 数据佐证（下钻 + 9 因素）</summary>')
        out.append(drill_h)
        out.append(factor_h)
        out.append('</details>')
    out.append('</div>')
    return ''.join(out)


def render_html(J):
    m = J['_meta']
    city = m.get('city', '全省')
    P = []
    P.append('<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">')
    P.append('<meta name="viewport" content="width=device-width,initial-scale=1.0">')
    P.append(f'<title>产品归因 · {h(city)} · {m.get("period_start")}~{m.get("period_end")}</title>')
    P.append(f'<style>{CSS}{EXTRA_CSS}</style></head><body>')
    P.append('<div class="toolbar"><button id="viewBtn" onclick="toggleMobile()">📱 手机版</button>'
             '<button id="exportBtn" onclick="exportLongImage()">📷 导出长图</button></div>')
    P.append('<header><div class="inner"><h1>🧩 产品视角经营归因</h1>')
    P.append(f'<div class="sub">{h(city)} · 本期 {m.get("period_start")}~{m.get("period_end")} · '
             f'同比 {m.get("yoy_start")}~{m.get("yoy_end")} / 环比 {m.get("mom_start")}~{m.get("mom_end")} · '
             f'好坏按同比增量金额排 · 生成 {m.get("generated_at")}</div>')
    P.append('</div></header><div class="wrap">')

    ov = J.get('总览归因', {}) or {}
    if ov.get('综述'):
        P.append('<h2>① 总览</h2>')
        c = analysis_html(ov['综述'], css_class='card', prefix='')
        c = c.replace('<p>', '<p><b class="hl">核心结论:</b>', 1)
        P.append(c)
    if ov.get('主线'):
        P.append('<div class="callout"><b>🔗 结构主线</b>'
                 + ''.join(f'<p>{h(x)}</p>' for x in ov['主线']) + '</div>')

    bg = _background_html(J.get('经营背景', {}))
    if bg:
        P.append('<h2>② 经营背景速览（当期 vs 同期）</h2>')
        P.append(bg)

    for tag, title in [('亮点子系列', '③ 🟢 亮点子系列（同比增量贡献最大）'),
                       ('问题子系列', '④ 🔴 问题子系列（同比下滑最大）')]:
        items = J.get(tag, [])
        if items:
            P.append(f'<h2>{title}</h2>')
            for s in items:
                P.append(_slice_card(s, tag))

    P.append('<footer>产品视角经营归因 · 数据由 SQL 算准(全量感知 SO/万,好坏按同比增量金额排)、'
             '归因由服务器 AI 撰写 · 数据源:product_flow_v / install_redpack_v / distribution_info / '
             'inventory_snapshot / promotion_meeting_summary / visit_record / dahua_redpack_quota / '
             'provider_contract / provider_storefront_invest</footer>')
    P.append('</div>')
    fbase = json.dumps(f"产品归因-{city}-{m.get('period_start')}_{m.get('period_end')}", ensure_ascii=False)
    P.append('<script src="https://cdn.jsdelivr.net/npm/html2canvas@1.4.1/dist/html2canvas.min.js"></script>')
    P.append('<script>')
    P.append('function toggleMobile(){var m=document.body.classList.toggle("mobile");'
             'document.getElementById("viewBtn").textContent=m?"🖥 桌面版":"📱 手机版";}')
    P.append('async function exportLongImage(){var btn=document.getElementById("exportBtn");'
             'if(typeof html2canvas==="undefined"){alert("长图组件未加载,请联网后刷新重试");return;}'
             'var mobile=document.body.classList.contains("mobile");'
             'btn.disabled=true;btn.textContent="生成中…";document.body.classList.add("exporting");'
             'await new Promise(function(r){setTimeout(r,60);});try{'
             'var canvas=await html2canvas(document.body,{scale:2,useCORS:true,backgroundColor:"#f4f6fa",'
             'width:document.body.clientWidth,windowWidth:document.body.clientWidth});'
             'var a=document.createElement("a");a.download=' + fbase + '+(mobile?"-手机版":"")+".png";'
             'a.href=canvas.toDataURL("image/png");a.click();}'
             'catch(e){alert("导出失败:"+e.message);}'
             'finally{document.body.classList.remove("exporting");btn.disabled=false;btn.textContent="📷 导出长图";}}')
    P.append('</script></body></html>')
    return '\n'.join(P)


# ────────────────────────── docx ──────────────────────────
def render_docx(J, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    m = J['_meta']
    city = m.get('city', '全省')
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(10.5)
    try:
        style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass

    def heading(txt, level=1):
        p = doc.add_paragraph()
        run = p.add_run(txt)
        run.bold = True
        run.font.size = Pt(15 - level * 1.5)
        run.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
        try:
            run.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        except Exception:
            pass
        return p

    def para(txt, **kw):
        return _para(doc, qn, Pt, RGBColor, txt, **kw)

    def analysis(value, color=(0x1f, 0x9d, 0x55), prefix='🔍 '):
        _ol_or_p_docx(doc, qn, Pt, RGBColor, value, color=color, prefix=prefix)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = 'Light Grid Accent 1'
        for k, hd in enumerate(headers):
            c = t.rows[0].cells[k]
            c.text = str(hd)
            for r in c.paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(9)
        for row in rows:
            cells = t.add_row().cells
            for k, v in enumerate(row):
                cells[k].text = '' if v is None else str(v)
                for r in cells[k].paragraphs[0].runs:
                    r.font.size = Pt(9)
        return t

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(f'产品视角经营归因 · {city} · {m.get("period_start")}~{m.get("period_end")}')
    tr.bold = True
    tr.font.size = Pt(18)
    tr.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
    try:
        tr.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass
    para(f'同比 {m.get("yoy_start")}~{m.get("yoy_end")} / 环比 {m.get("mom_start")}~{m.get("mom_end")} · '
         f'好坏按同比增量金额排 · 生成 {m.get("generated_at")}', size=9, color=(0x6b, 0x76, 0x86))

    ov = J.get('总览归因', {}) or {}
    if ov.get('综述'):
        heading('① 总览', 1)
        analysis(ov['综述'], color=None, prefix='')
    if ov.get('主线'):
        heading('结构主线', 2)
        analysis(ov['主线'], color=(0x1f, 0x5f, 0xbf), prefix='🔗 ')

    bg = J.get('经营背景', {}) or {}
    if bg:
        heading('② 经营背景速览', 1)
        pm = (bg.get('推广会') or {}).get('当期')
        if pm:
            para(f"推广会:{i(pm.get('场次'))} 场 / 签到 {i(pm.get('签到公司'))} 家,"
                 f"会后红包 {i(pm.get('会后红包金额'))} 元 / 夜视王上线 {i(pm.get('夜视王上线台数'))} 台")
        vs = (bg.get('跑动') or {}).get('当期')
        if vs:
            para(f"跑动:{i(vs.get('拜访次数'))} 次 / 覆盖 {i(vs.get('覆盖服务商'))} 家,大华打卡 {i(vs.get('大华打卡次数'))} 次")
        sr = (bg.get('业务员红包') or {}).get('当期')
        if sr:
            para(f"业务员红包:使用 {i(sr.get('红包使用金额'))} 元 / {i(sr.get('发放客户数'))} 客户,"
                 f"转化激活 {i(sr.get('转化激活客户'))} 户,解锁率 {pct(sr.get('平均解锁率'))}")
        sf = bg.get('门头') or {}
        if sf.get('投入服务商') is not None:
            para(f"门头({sf.get('年份')}年):投入 {i(sf.get('投入服务商'))} 家 / {wan(sf.get('投入金额_万'))} 万,"
                 f"激活 {i(sf.get('激活数'))} 家,带来上线 {wan(sf.get('带来上线金额_万'))} 万")

    for tag, title2 in [('亮点子系列', '③ 亮点子系列（增量贡献最大）'),
                        ('问题子系列', '④ 问题子系列（下滑最大）')]:
        items = J.get(tag, [])
        if not items:
            continue
        heading(title2, 1)
        for s in items:
            g = s.get('归因') or {}
            heading(f"{s.get('子系列') or '(未命名子系列)'}　{swan(s.get('增量_万'))}万 · "
                    f"同比{pct(s.get('同比'))} · 货值{wan(s.get('货值_万'))}万", 2)
            if g.get('定性'):
                para('判定:' + g['定性'], bold=True)
            if g.get('归因归咎'):
                analysis(g['归因归咎'], prefix='')
            if g.get('关注点'):
                analysis(g['关注点'], color=(0xd9, 0x8a, 0x00), prefix='👉 ')
            # 佐证:地市 + 代理商下钻表
            for dim in ['地市', '代理商']:
                d = (s.get('下钻') or {}).get(dim, {}) or {}
                rows = _drill_rows(d, dim)
                if rows:
                    para(f'{dim}涨跌:', size=9, bold=True)
                    table([dim, '增量万', '货值万'], [[r[0], r[1][0] if isinstance(r[1], tuple) else r[1], r[2]] for r in rows])

    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--in', dest='inp', required=True, help='attribution.json(数据+AI归因)')
    ap.add_argument('--outdir', required=True)
    ap.add_argument('--basename', default=None)
    ap.add_argument('--no-docx', action='store_true')
    args = ap.parse_args()

    J = json.load(open(args.inp, encoding='utf-8'))
    m = J['_meta']
    base = args.basename or f"产品归因-{m.get('city', '全省')}-{m.get('period_start')}_{m.get('period_end')}"
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    html_path = outdir / f"{base}.html"
    html_path.write_text(render_html(J), encoding='utf-8')
    print(f"✓ HTML: {html_path}")

    if not args.no_docx:
        docx_path = outdir / f"{base}.docx"
        render_docx(J, str(docx_path))
        print(f"✓ docx: {docx_path}")


if __name__ == '__main__':
    main()
