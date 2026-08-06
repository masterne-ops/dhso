#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""全省周报渲染 —— 复用 render.py 助手 + render_city.py 的 data_table/手机版能力。

输入 gather_prov_weekly 输出的 JSON（+ LLM 写的「分析」字段）。
产出 HTML(含手机版切换+导出长图) + docx。

数据键对应 gather_prov_weekly：①周趋势 / ②地市对比 / ③异动下钻 / ④专项 / 分析。
"""
import argparse
import json
from pathlib import Path

from render import pct, cls, h, analysis_html, CSS
from render_city import data_table, _para, _ol_or_p_docx


def _rate(x):
    return f'{x*100:.0f}%' if x is not None else 'N/A'


def render_html(J):
    T = J['①周趋势']
    C = J['②地市对比']
    Y = J.get('②YTD累计同比', {})
    PBC = J.get('服务商按地市', {})
    PK = J.get('服务商KPI', {})
    ACT = J.get('激活V2', {})
    VS = J.get('跑动分组', {})
    D = J['③异动下钻']
    F = J['④专项']
    A = J.get('分析', {})
    WK = J['本周区间']
    wins = F.get('标签', T.get('标签', []))

    P = []
    P.append('<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">')
    P.append('<meta name="viewport" content="width=device-width,initial-scale=1.0">')
    P.append(f'<title>浙江全省周报 · {WK}</title><style>{CSS}')
    P.append('.toolbar{position:fixed;top:14px;right:14px;z-index:999;display:flex;gap:8px}'
             '.toolbar button{background:#fff;color:var(--pri-d);border:1px solid var(--pri);border-radius:8px;'
             'padding:7px 12px;font-size:13px;font-weight:600;cursor:pointer;box-shadow:0 2px 8px rgba(20,40,80,.15)}'
             '.toolbar button:hover{background:var(--pri);color:#fff}.toolbar button:disabled{opacity:.6;cursor:wait}'
             '.exporting .toolbar{display:none!important}@media print{.toolbar{display:none!important}}')
    P.append('.cards{display:none}.dcard{border:1px solid var(--line);border-radius:9px;padding:9px 11px;margin:7px 0;background:#fafbfd}'
             '.dcard .nm{font-weight:700;font-size:13.5px;margin-bottom:5px}'
             '.dcard .kv{display:flex;justify-content:space-between;font-size:12.5px;padding:2px 0;border-top:1px dashed var(--line)}'
             '.dcard .kv:first-of-type{border-top:none}.dcard .kv .k{color:var(--mute)}.dcard .kv .v{font-variant-numeric:tabular-nums;font-weight:600}')
    P.append('body.mobile .wrap{max-width:440px}body.mobile header .inner{max-width:440px}'
             'body.mobile .tbl{display:none}body.mobile .cards{display:block}')
    P.append('@media(max-width:520px){.wrap{max-width:100%;padding:0 12px 60px}.tbl{display:none}.cards{display:block}h1{font-size:20px}}')
    P.append('.bar{background:#eef1f6;border-radius:5px;height:16px;position:relative;overflow:hidden;margin:3px 0}'
             '.bar>i{display:block;height:100%;background:linear-gradient(90deg,#2f80ed,#1f5fbf)}'
             '.bar>.t{position:absolute;top:-1px;height:18px;border-left:2px dashed #d64545}')
    P.append('</style></head><body>')
    P.append('<div class="toolbar"><button id="viewBtn" onclick="toggleMobile()">📱 手机版</button>'
             '<button id="exportBtn" onclick="exportLongImage()">📷 导出长图</button></div>')
    P.append('<header><div class="inner"><h1>浙江全省经营周报</h1>')
    P.append(f'<div class="sub">本周 {WK}（7天窗口）· 主对比：周环比上周 · 口径：全量感知 SO（万），全省=浙江11地市之和</div>')
    P.append('</div></header><div class="wrap">')

    # 核心结论
    if A.get('总览'):
        ov = analysis_html(A['总览'], css_class='card', prefix='')
        ov = ov.replace('<p>', '<p><b class="hl">核心结论：</b>', 1)
        P.append(ov)

    # ① 全省周趋势
    P.append('<h2>① 全省 SO 周趋势</h2><div class="card">')
    trend_cells = '　'.join(f"{lab} <b class='num'>{so}</b>"
                            for lab, so in zip(T['标签'], T['分周SO']))
    P.append(f"<p>最近 {len(T['标签'])} 周（7天窗口，万元）：{trend_cells}</p>")
    P.append(f"<p>本周 <b class='num'>{T['本周SO']}</b> 万，环比上周 "
             f"<span class='{cls(T['本周环比'])}'>{pct(T['本周环比'])}</span></p>")
    # 全省当月完成率进度条(参考地市周报)：当月累计=②YTD全省当月今年;应达成=各市应达成之和
    _acc = Y.get('全省当月今年') or 0
    _need = round(sum((r.get('当月应达成') or 0) for r in C['地市']), 1)
    _rt = (_acc / _need) if _need else 0
    _tp = C.get('当月时间进度') or 0
    P.append(f"<h3>全省当月完成率（{C.get('当月', '')}）</h3>")
    P.append(f"<div class='bar'><i style='width:{min(_rt*100,100):.0f}%'></i>"
             f"<span class='t' style='left:{min(_tp*100,100):.0f}%'></span></div>")
    P.append(f"<p class='mut'>当月累计 {_acc} 万 / 应达成 {_need} 万，完成率 <b>{_rt*100:.0f}%</b>；"
             f"时间进度 {_tp*100:.0f}%（红虚线）——完成率达时间进度即为跟上节奏。</p>")
    P.append(f"<p>📊 <b>当月同比</b>：本月至今 <b class='num'>{Y.get('全省当月今年')}</b> 万，较去年同期 {Y.get('全省当月去年')} 万 "
             f"<span class='{cls(Y.get('全省当月同比'))}'>{pct(Y.get('全省当月同比'))}</span>。</p>")
    P.append(f"<p>📈 <b>YTD 累计同比</b>：今年至今 <b class='num'>{Y.get('全省今年')}</b> 万，较去年同期 {Y.get('全省去年')} 万 "
             f"<span class='{cls(Y.get('全省累计同比'))}'>{pct(Y.get('全省累计同比'))}</span>。</p></div>")

    # ② 11地市周对比
    P.append('<h2>② 11 地市周对比</h2><div class="card">')
    P.append(data_table(
        ['地市', '本周SO万', '上周SO万', '本周环比', '当月完成率'],
        [[h(r['市']), r['本周SO'], r['上周SO'], (pct(r['本周环比']), cls(r['本周环比'])),
          pct(r.get('当月完成率'))] for r in C['地市']]))
    rb = '；'.join(f"{x['市']}（{pct(x['本周环比'])}）" for x in C['红榜'])
    hb = '；'.join(f"{x['市']}（{pct(x['本周环比'])}）" for x in C['黑榜'])
    P.append(f"<p>🏆 <b>本周红榜</b>：{rb}</p><p>⚠️ <b>本周黑榜</b>：{hb}</p>")
    if A.get('地市'):
        P.append(analysis_html(A['地市']))
    P.append('</div>')

    # ③ SO 同比（当月 + YTD）
    if Y.get('地市'):
        P.append('<h2>③ SO 同比（当月当天累计 + 年初累计）</h2><div class="card">')
        P.append(f"<p>全省当月（{Y.get('当月区间','')} vs 去年同期）：<b class='num'>{Y.get('全省当月今年')}</b> 万，"
                 f"较去年同期 <span class='{cls(Y.get('全省当月同比'))}'>{pct(Y.get('全省当月同比'))}</span>；"
                 f"全省 YTD：<b class='num'>{Y.get('全省今年')}</b> 万，"
                 f"较去年 <span class='{cls(Y.get('全省累计同比'))}'>{pct(Y.get('全省累计同比'))}</span>。</p>")
        P.append(data_table(
            ['地市', '当月今年万', '当月去年万', '当月同比', 'YTD今年万', 'YTD去年万', '累计同比'],
            [[h(r['市']), r['当月今年'], r['当月去年'], (pct(r['当月同比']), cls(r['当月同比'])),
              r['YTD今年'], r['YTD去年'], (pct(r['累计同比']), cls(r['累计同比']))] for r in Y['地市']]))
        P.append('</div>')

    # ④ 服务商发展（按11地市）+ 跑动分组
    if PBC.get('地市') or PK:
        P.append('<h2>④ 服务商发展（本周签约/激活 + 任务达成）</h2><div class="card">')
        if PK:
            actcur = ACT.get('本周达V2', PK.get('本周激活'))
            P.append(f"<p>全省本周新签约 <b>{PK['本周签约']}</b> 家；本周激活（官方激活时间落本周）<b>{actcur}</b> 家"
                     f"，环比 <span class='{cls(ACT.get('环比'))}'>{pct(ACT.get('环比'))}</span>。</p>")
        if PBC.get('地市'):
            P.append('<h3>11 地市服务商发展对比</h3>')
            P.append(data_table(
                ['地市', '本周签约', '签约达成率', '本周激活(达V2)', '激活环比', '激活达成率'],
                [[h(r['市']), r['本周签约'], _rate(r['签约达成率']), r['本周激活'],
                  (pct(r['激活环比']), cls(r['激活环比'])), _rate(r['激活达成率'])] for r in PBC['地市']]))
        if PK.get('任务达成'):
            td = PK['任务达成']
            P.append('<h3>全省任务达成（当前实际 / 年度目标）</h3>')
            P.append(data_table(['指标', '当前实际', '年度目标', '达成率'],
                     [[k, td[k]['实际'], td[k]['目标'], (_rate(td[k]['达成率']),
                       cls((td[k]['达成率'] or 0) - 0.999))] for k in ['签约数', '新签', 'V2', 'V3', 'V4']]))
        if VS:
            dh, ag = VS.get('大华', {}), VS.get('代理商', {})
            P.append('<h3>跑动人员分类（大华 / 代理商业务员）</h3>')
            P.append(data_table(['人员', '本周打卡', '本周人数', '上周打卡', '打卡环比'],
                     [['🏢 大华业务员', dh.get('本周打卡', 0), dh.get('本周人数', 0), dh.get('上周打卡', 0),
                       (pct(dh.get('打卡环比')), cls(dh.get('打卡环比')))],
                      ['🏪 代理商业务员', ag.get('本周打卡', 0), ag.get('本周人数', 0), ag.get('上周打卡', 0),
                       (pct(ag.get('打卡环比')), cls(ag.get('打卡环比')))]]))
        if A.get('服务商'):
            P.append(analysis_html(A['服务商']))
        P.append('</div>')

    # ⑤ 异动地市下钻
    if D.get('明细'):
        P.append('<h2>⑤ 异动地市下钻</h2><div class="card">')
        if A.get('下钻'):
            P.append(analysis_html(A['下钻']))
        for dd in D['明细']:
            rd = dd['跑动']
            P.append(f"<h3>🔴 {h(dd['地市'])} <span class='mut'>（周跑动打卡 "
                     f"{rd['上周打卡']}→{rd['本周打卡']}，环比 {pct(rd['环比'])}）</span></h3>")
            if dd['问题代理商']:
                P.append(data_table(
                    ['代理商', '本周出货万', '上周出货万', '环比'],
                    [[h(x['代理商']), x['本周'], x['上周'], (pct(x['环比']), cls(x['环比']))]
                     for x in dd['问题代理商']]))
            else:
                P.append("<p class='mut'>本周无明显异动代理商。</p>")
        P.append('</div>')

    # ⑥ 三专项周趋势
    P.append('<h2>⑥ 三大产品专项周趋势（台数）</h2><div class="card">')
    P.append(data_table(
        ['专项'] + wins + ['本周环比', '状态'],
        [[x['专项']] + x['分周台数'] +
         [(pct(x['本周环比']), cls(x['本周环比'])),
          ('<span class="pill p-red">下降</span>' if x['是否下降'] else '<span class="pill p-green">增长</span>')]
         for x in F['专项']]))
    if A.get('专项'):
        P.append(analysis_html(A['专项']))
    P.append('</div>')

    # ⑦ NP 转入客户进展（11地市对比）
    NPT = J.get('NP转入')
    if NPT and NPT.get('地市'):
        P.append('<h2>⑦ NP 转入客户进展（11 地市）</h2><div class="card">')
        _nprows = [[h(r['市']), r['转入'], r['已报备'], r['相关'], r['已签约'], r['已激活'], pct(r['签约率'])]
                   for r in NPT['地市']]
        _npt = NPT['合计']
        _nprows.append(['全省合计', _npt['转入'], _npt['已报备'], _npt['相关'], _npt['已签约'], _npt['已激活'], pct(_npt['签约率'])])
        P.append(data_table(['地市', '转入', '已报备', '相关', '已签约', '已激活', '签约率'], _nprows))
        P.append(f"<p class='mut'>NP 转入客户转化漏斗，截至最新快照（{NPT.get('数据时点', '')}）当前状态；"
                 "已签约=命中签约表，已激活=累计上线≥1000 元。旧批「NP 流转名单」无报备/相关字段，报备/相关仅反映新批。</p>")
        P.append('</div>')

    # ⑧ 下阶段计划
    plans = A.get('下阶段计划', [])
    if plans:
        P.append('<h2>⑧ 下阶段计划</h2>')
        P.append(analysis_html(plans, css_class='card', prefix=''))

    P.append('<footer>浙江全省经营周报 · 数据源：生产库 product_flow_v / visit_record_v / product_focus · '
             '7天窗口 · 数字由 report_tools 计算，分析由 AI 撰写</footer>')
    P.append('</div>')
    fbase = json.dumps(f"浙江全省周报-{WK}", ensure_ascii=False)
    P.append('<script src="https://cdn.jsdelivr.net/npm/html2canvas@1.4.1/dist/html2canvas.min.js"></script>')
    P.append('<script>')
    P.append('function toggleMobile(){var m=document.body.classList.toggle("mobile");'
             'document.getElementById("viewBtn").textContent=m?"🖥 桌面版":"📱 手机版";}')
    P.append('async function exportLongImage(){var btn=document.getElementById("exportBtn");'
             'if(typeof html2canvas==="undefined"){alert("长图组件未加载，请联网后刷新重试");return;}'
             'var mobile=document.body.classList.contains("mobile");'
             'btn.disabled=true;btn.textContent="生成中…";document.body.classList.add("exporting");'
             'await new Promise(function(r){setTimeout(r,60);});'
             'try{var canvas=await html2canvas(document.body,{scale:2,useCORS:true,backgroundColor:"#f4f6fa",'
             'width:document.body.clientWidth,windowWidth:document.body.clientWidth});'
             'var a=document.createElement("a");a.download=' + fbase + '+(mobile?"-手机版":"")+".png";'
             'a.href=canvas.toDataURL("image/png");a.click();}'
             'catch(e){alert("导出失败："+e.message);}'
             'finally{document.body.classList.remove("exporting");btn.disabled=false;btn.textContent="📷 导出长图";}}')
    P.append('</script></body></html>')
    return '\n'.join(P)


def render_docx(J, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    T = J['①周趋势']; C = J['②地市对比']
    Y = J.get('②YTD累计同比', {}); PBC = J.get('服务商按地市', {})
    PK = J.get('服务商KPI', {}); ACT = J.get('激活V2', {}); VS = J.get('跑动分组', {})
    D = J['③异动下钻']; F = J['④专项']
    A = J.get('分析', {}); WK = J['本周区间']
    wins = F.get('标签', T.get('标签', []))

    doc = Document()
    st = doc.styles['Normal']; st.font.name = '宋体'; st.font.size = Pt(10.5)
    try: st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception: pass

    def heading(txt, lv=1):
        p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True
        r.font.size = Pt(15 - lv * 1.5); r.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
        try: r.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        except Exception: pass

    def para(txt, **kw): return _para(doc, qn, Pt, RGBColor, txt, **kw)
    def analysis(v, color=(0x1f, 0x9d, 0x55), prefix='🔍 '): _ol_or_p_docx(doc, qn, Pt, RGBColor, v, color=color, prefix=prefix)

    def table(headers, rows):
        tb = doc.add_table(rows=1, cols=len(headers)); tb.style = 'Light Grid Accent 1'
        for i, hd in enumerate(headers):
            c = tb.rows[0].cells[i]; c.text = str(hd)
            for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(9)
        for row in rows:
            cs = tb.add_row().cells
            for i, v in enumerate(row):
                cs[i].text = '' if v is None else str(v)
                for r in cs[i].paragraphs[0].runs: r.font.size = Pt(9)

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(f'浙江全省经营周报 · {WK}'); tr.bold = True; tr.font.size = Pt(18)
    tr.font.color.rgb = RGBColor(0x14, 0x47, 0x8f)
    try: tr.font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception: pass
    para('本周7天窗口 · 主对比周环比上周 · 口径：全量感知SO（万），全省=11地市之和', size=9, color=(0x6b, 0x76, 0x86))

    if A.get('总览'): heading('核心结论', 2); analysis(A['总览'], color=None, prefix='')
    heading('① 全省 SO 周趋势', 1)
    para('最近 %d 周：' % len(T['标签']) + '　'.join(f"{lab} {so}" for lab, so in zip(T['标签'], T['分周SO'])))
    para(f"本周 {T['本周SO']} 万，环比上周 {pct(T['本周环比'])}", bold=True)
    _acc = Y.get('全省当月今年') or 0
    _need = round(sum((r.get('当月应达成') or 0) for r in C['地市']), 1)
    _rt = (_acc / _need) if _need else 0
    _tp = C.get('当月时间进度') or 0
    para(f"全省当月完成率（{C.get('当月', '')}）：当月累计 {_acc} 万 / 应达成 {_need} 万 = {_rt*100:.0f}%；"
         f"时间进度 {_tp*100:.0f}%（完成率达时间进度即为跟上节奏）。", bold=True)
    para(f"当月同比：本月至今 {Y.get('全省当月今年')} 万，较去年同期 {Y.get('全省当月去年')} 万 {pct(Y.get('全省当月同比'))}；"
         f"YTD 累计同比：今年至今 {Y.get('全省今年')} 万，较去年 {Y.get('全省去年')} 万 {pct(Y.get('全省累计同比'))}。", bold=True)
    heading('② 11 地市周对比', 1)
    table(['地市', '本周SO万', '上周SO万', '本周环比', '当月完成率'],
          [[r['市'], r['本周SO'], r['上周SO'], pct(r['本周环比']), pct(r.get('当月完成率'))] for r in C['地市']])
    para('🏆 红榜：' + '；'.join(f"{x['市']}（{pct(x['本周环比'])}）" for x in C['红榜']))
    para('⚠️ 黑榜：' + '；'.join(f"{x['市']}（{pct(x['本周环比'])}）" for x in C['黑榜']))
    if A.get('地市'): analysis(A['地市'])
    # ③ SO 同比
    if Y.get('地市'):
        heading('③ SO 同比（当月当天累计 + 年初累计）', 1)
        para(f"全省当月 {Y.get('全省当月今年')} 万（{pct(Y.get('全省当月同比'))}）；"
             f"全省 YTD {Y.get('全省今年')} 万（{pct(Y.get('全省累计同比'))}）。", bold=True)
        table(['地市', '当月今年万', '当月去年万', '当月同比', 'YTD今年万', 'YTD去年万', '累计同比'],
              [[r['市'], r['当月今年'], r['当月去年'], pct(r['当月同比']),
                r['YTD今年'], r['YTD去年'], pct(r['累计同比'])] for r in Y['地市']])
    # ④ 服务商发展（按地市）+ 跑动分组
    if PBC.get('地市') or PK:
        heading('④ 服务商发展（本周签约/激活 + 任务达成）', 1)
        if PK:
            actcur = ACT.get('本周达V2', PK.get('本周激活'))
            para(f"全省本周新签约 {PK['本周签约']} 家；本周激活（官方激活时间落本周）{actcur} 家，环比 {pct(ACT.get('环比'))}。", bold=True)
        if PBC.get('地市'):
            para('11 地市服务商发展对比：')
            table(['地市', '本周签约', '签约达成率', '本周激活(达V2)', '激活环比', '激活达成率'],
                  [[r['市'], r['本周签约'], _rate(r['签约达成率']), r['本周激活'],
                    pct(r['激活环比']), _rate(r['激活达成率'])] for r in PBC['地市']])
        if PK.get('任务达成'):
            td = PK['任务达成']
            para('全省任务达成（当前实际 / 年度目标）：')
            table(['指标', '当前实际', '年度目标', '达成率'],
                  [[k, td[k]['实际'], td[k]['目标'], _rate(td[k]['达成率'])] for k in ['签约数', '新签', 'V2', 'V3', 'V4']])
        if VS:
            dh, ag = VS.get('大华', {}), VS.get('代理商', {})
            para('跑动人员分类（大华 / 代理商业务员）：', bold=True)
            table(['人员', '本周打卡', '本周人数', '上周打卡', '打卡环比'],
                  [['大华业务员', dh.get('本周打卡', 0), dh.get('本周人数', 0), dh.get('上周打卡', 0), pct(dh.get('打卡环比'))],
                   ['代理商业务员', ag.get('本周打卡', 0), ag.get('本周人数', 0), ag.get('上周打卡', 0), pct(ag.get('打卡环比'))]])
        if A.get('服务商'): analysis(A['服务商'])
    if D.get('明细'):
        heading('⑤ 异动地市下钻', 1)
        if A.get('下钻'): analysis(A['下钻'])
        for dd in D['明细']:
            rd = dd['跑动']
            heading(f"🔴 {dd['地市']}（周跑动 {rd['上周打卡']}→{rd['本周打卡']}，环比 {pct(rd['环比'])}）", 2)
            if dd['问题代理商']:
                table(['代理商', '本周出货万', '上周出货万', '环比'],
                      [[x['代理商'], x['本周'], x['上周'], pct(x['环比'])] for x in dd['问题代理商']])
            else:
                para('本周无明显异动代理商。', size=9, color=(0x6b, 0x76, 0x86))
    heading('⑥ 三大产品专项周趋势（台数）', 1)
    table(['专项'] + wins + ['本周环比', '状态'],
          [[x['专项']] + x['分周台数'] + [pct(x['本周环比']), '下降' if x['是否下降'] else '增长'] for x in F['专项']])
    if A.get('专项'): analysis(A['专项'])
    NPT = J.get('NP转入')
    if NPT and NPT.get('地市'):
        heading('⑦ NP 转入客户进展（11 地市）', 1)
        _nprows = [[r['市'], r['转入'], r['已报备'], r['相关'], r['已签约'], r['已激活'], pct(r['签约率'])]
                   for r in NPT['地市']]
        _npt = NPT['合计']
        _nprows.append(['全省合计', _npt['转入'], _npt['已报备'], _npt['相关'], _npt['已签约'], _npt['已激活'], pct(_npt['签约率'])])
        table(['地市', '转入', '已报备', '相关', '已签约', '已激活', '签约率'], _nprows)
        para('NP 转入客户转化漏斗，截至最新快照当前状态；已签约=命中签约表，已激活=累计上线≥1000元。旧批无报备/相关字段。',
             size=9, color=(0x6b, 0x76, 0x86))
    plans = A.get('下阶段计划', [])
    if plans: heading('⑧ 下阶段计划', 1); analysis(plans, color=None, prefix='')
    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--in', dest='inp', required=True)
    ap.add_argument('--outdir', required=True)
    ap.add_argument('--basename', default=None)
    ap.add_argument('--no-docx', action='store_true')
    args = ap.parse_args()
    J = json.load(open(args.inp, encoding='utf-8'))
    base = args.basename or f"浙江全省周报-{J['本周区间']}"
    outdir = Path(args.outdir); outdir.mkdir(parents=True, exist_ok=True)
    (outdir / f"{base}.html").write_text(render_html(J), encoding='utf-8')
    print(f"✓ HTML: {outdir / (base + '.html')}")
    if not args.no_docx:
        render_docx(J, str(outdir / f"{base}.docx"))
        print(f"✓ docx: {outdir / (base + '.docx')}")


if __name__ == '__main__':
    main()
