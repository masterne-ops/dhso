#!/usr/bin/env python3
"""高德潜客 + NP 流转名单入库 + 匹配校验

业务理解:
  - 潜客 / NP 流转 = 还没登记到系统的服务商
  - 高德潜客「客户名称」非空 = 已签约(我方录入内部名)
  - NP 流转「签约服务商名称」非空 = 已签约(到 provider_contract 反查验证)
"""
import sqlite3
import openpyxl
from datetime import datetime

DB = '/opt/so-data-analytics/db/product_flow.db'


# ─── DDL ───
DDL = """
CREATE TABLE IF NOT EXISTS gaode_potential_customer (
    数据时点              TEXT NOT NULL,
    外部客户名称          TEXT NOT NULL,
    客户来源              TEXT,
    客户编码              TEXT,
    客户名称              TEXT,            -- 非 NULL = 已签约
    是否已转出            TEXT,
    与我司业务相关        TEXT,
    下一步计划            TEXT,
    原因说明              TEXT,
    预计完成时间          TEXT,
    是否转出组织          TEXT,
    签约类型              TEXT,
    客户状态              TEXT,
    客户所有者工号        TEXT,
    客户所有者            TEXT,
    客户所有者部门        TEXT,
    年安防采购量          TEXT,
    对大华品牌认可度      TEXT,
    责任分销商            TEXT,
    责任分销商分配人      TEXT,
    责任分销商分配时间    TEXT,
    责任分销经理          TEXT,
    责任分销经理工号      TEXT,
    责任分销经理分配人    TEXT,
    责任分销经理分配时间  TEXT,
    客户联系人电话        TEXT,
    客户联系地址          TEXT,
    省                    TEXT,
    市                    TEXT,
    区县                  TEXT,
    分销商认证            TEXT,
    渠道客户类型          TEXT,
    协同人姓名            TEXT,
    安装红包上线金额      REAL,
    最近拜访时间_分销经理 TEXT,
    年度拜访次数_分销经理 INTEGER,
    本月拜访次数_分销经理 INTEGER,
    最近拜访时间_分销商   TEXT,
    年度拜访次数_分销商   INTEGER,
    本月拜访次数_分销商   INTEGER,
    创建时间              TEXT,
    外部客户关联时间      TEXT,
    客户关闭时间          TEXT,
    转公共池时间          TEXT,
    签约时间              TEXT,
    首次梳理时间          TEXT,
    最后1次梳理时间       TEXT,
    PRIMARY KEY (数据时点, 外部客户名称)
);
CREATE INDEX IF NOT EXISTS idx_gpc_city ON gaode_potential_customer(市);
CREATE INDEX IF NOT EXISTS idx_gpc_district ON gaode_potential_customer(区县);
CREATE INDEX IF NOT EXISTS idx_gpc_mgr ON gaode_potential_customer(责任分销经理);
CREATE INDEX IF NOT EXISTS idx_gpc_signed ON gaode_potential_customer(客户名称);

CREATE TABLE IF NOT EXISTS np_customer_pool (
    数据时点         TEXT NOT NULL,
    客户名称         TEXT NOT NULL,
    地市             TEXT,
    区县             TEXT,
    区县所有者负责人 TEXT,
    签约服务商名称   TEXT,            -- 非 NULL = 已签约
    PRIMARY KEY (数据时点, 客户名称)
);
CREATE INDEX IF NOT EXISTS idx_npc_city ON np_customer_pool(地市);
CREATE INDEX IF NOT EXISTS idx_npc_signed ON np_customer_pool(签约服务商名称);
"""


def to_cell(v):
    if v is None: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
    return v


def load_gaode(conn, fp, period):
    print(f'\n══ 高德潜客 {fp} ──> 数据时点 {period} ══')
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = wb['sheet1']
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip().replace('\n', '_') if c else None for c in rows[0]]
    print(f'  Excel: {len(rows)-1} 行,{len(hdr)} 列')

    # Excel 列名 → DB 列名(部分需重命名)
    rename = {
        '最近拜访时间_（分销经理）': '最近拜访时间_分销经理',
        '年度拜访次数_（分销经理）': '年度拜访次数_分销经理',
        '本月拜访次数_（分销经理）': '本月拜访次数_分销经理',
        '最近拜访时间_（分销商）':   '最近拜访时间_分销商',
        '年度拜访次数_（分销商）':   '年度拜访次数_分销商',
        '本月拜访次数_（分销商）':   '本月拜访次数_分销商',
    }
    db_cols = [r[1] for r in conn.execute('PRAGMA table_info(gaode_potential_customer)').fetchall()]
    col_map = []  # [(excel_idx, db_col)]
    for i, h in enumerate(hdr):
        if not h: continue
        db_h = rename.get(h, h)
        if db_h in db_cols:
            col_map.append((i, db_h))
        else:
            print(f'  ⚠️ 跳过未映射列: {repr(h)}')

    print(f'  对齐: {len(col_map)} 列')
    conn.execute("DELETE FROM gaode_potential_customer WHERE 数据时点 = ?", (period,))

    col_sql = ', '.join(f'"{c}"' for c in ['数据时点'] + [c for _, c in col_map])
    ph = ', '.join(['?'] * (1 + len(col_map)))
    insert_sql = f'INSERT OR IGNORE INTO gaode_potential_customer ({col_sql}) VALUES ({ph})'

    batch = []
    for r in rows[1:]:
        if not r or not any(r): continue
        ext_name = r[3] if 3 < len(r) else None  # 外部客户名称是 PK 一部分
        if not ext_name: continue
        vals = [period] + [to_cell(r[i]) if i < len(r) else None for i, _ in col_map]
        batch.append(tuple(vals))

    cur = conn.executemany(insert_sql, batch)
    conn.commit()
    n = conn.execute("SELECT COUNT(*) FROM gaode_potential_customer WHERE 数据时点 = ?", (period,)).fetchone()[0]
    print(f'  ✅ 入库 {cur.rowcount} 行 / 表中该时点 {n} 行')


def load_np(conn, fp, period):
    print(f'\n══ NP 流转 {fp} ──> 数据时点 {period} ══')
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = wb['Sheet3']
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else None for c in rows[0]]
    print(f'  Excel: {len(rows)-1} 行,字段 {hdr}')

    conn.execute("DELETE FROM np_customer_pool WHERE 数据时点 = ?", (period,))
    insert_sql = """INSERT OR IGNORE INTO np_customer_pool
        (数据时点, 客户名称, 地市, 区县, 区县所有者负责人, 签约服务商名称)
        VALUES (?, ?, ?, ?, ?, ?)"""
    batch = []
    for r in rows[1:]:
        if not r or not r[0]: continue
        batch.append((period, to_cell(r[0]), to_cell(r[1]), to_cell(r[2]), to_cell(r[3]), to_cell(r[4])))
    cur = conn.executemany(insert_sql, batch)
    conn.commit()
    n = conn.execute("SELECT COUNT(*) FROM np_customer_pool WHERE 数据时点 = ?", (period,)).fetchone()[0]
    print(f'  ✅ 入库 {cur.rowcount} 行 / 表中该时点 {n} 行')


def verify(conn):
    print('\n═══ 校验 ═══')
    # 高德
    print('\n— 高德潜客 —')
    total_g = conn.execute("SELECT COUNT(*) FROM gaode_potential_customer WHERE 数据时点 = '2026-05-16'").fetchone()[0]
    signed_g = conn.execute("SELECT COUNT(*) FROM gaode_potential_customer WHERE 数据时点 = '2026-05-16' AND 客户名称 IS NOT NULL AND 客户名称 != ''").fetchone()[0]
    matched_g = conn.execute("""
        SELECT COUNT(DISTINCT g.外部客户名称) FROM gaode_potential_customer g
         WHERE g.数据时点 = '2026-05-16'
           AND g.客户名称 IS NOT NULL AND g.客户名称 != ''
           AND EXISTS (SELECT 1 FROM provider_contract pc WHERE pc.客户名称 = g.客户名称)
    """).fetchone()[0]
    print(f'  累计下发: {total_g} 家')
    print(f'  已签约(客户名称非空): {signed_g} 家')
    print(f'  在 provider_contract 命中: {matched_g} 家')

    # 高德已激活(信号: install_redpack 累计金额 ≥ 1000)
    activated_g = conn.execute("""
        WITH signed AS (
          SELECT g.客户名称 FROM gaode_potential_customer g
           WHERE g.数据时点 = '2026-05-16' AND g.客户名称 IS NOT NULL AND g.客户名称 != ''
        ),
        amt AS (
          SELECT ir.上线客户名称, SUM(ir.产品现有分销价) AS 金额
            FROM install_redpack ir GROUP BY 1
        )
        SELECT COUNT(*) FROM signed s JOIN amt a ON a.上线客户名称 = s.客户名称
         WHERE a.金额 >= 1000
    """).fetchone()[0]
    print(f'  已激活(累计上线 ≥ 1000): {activated_g} 家')

    # NP
    print('\n— NP 流转 —')
    total_n = conn.execute("SELECT COUNT(*) FROM np_customer_pool WHERE 数据时点 = '2026-05-24'").fetchone()[0]
    signed_n = conn.execute("SELECT COUNT(*) FROM np_customer_pool WHERE 数据时点 = '2026-05-24' AND 签约服务商名称 IS NOT NULL AND 签约服务商名称 != ''").fetchone()[0]
    matched_n = conn.execute("""
        SELECT COUNT(DISTINCT n.客户名称) FROM np_customer_pool n
         WHERE n.数据时点 = '2026-05-24'
           AND n.签约服务商名称 IS NOT NULL AND n.签约服务商名称 != ''
           AND EXISTS (SELECT 1 FROM provider_contract pc WHERE pc.客户名称 = n.签约服务商名称)
    """).fetchone()[0]
    print(f'  识别有效流转: {total_n} 家')
    print(f'  已签约(签约服务商名称非空): {signed_n} 家')
    print(f'  在 provider_contract 命中: {matched_n} 家')

    activated_n = conn.execute("""
        WITH signed AS (
          SELECT n.签约服务商名称 AS 签约名 FROM np_customer_pool n
           WHERE n.数据时点 = '2026-05-24' AND n.签约服务商名称 IS NOT NULL AND n.签约服务商名称 != ''
        ),
        amt AS (
          SELECT ir.上线客户名称, SUM(ir.产品现有分销价) AS 金额
            FROM install_redpack ir GROUP BY 1
        )
        SELECT COUNT(*) FROM signed s JOIN amt a ON a.上线客户名称 = s.签约名 WHERE a.金额 >= 1000
    """).fetchone()[0]
    print(f'  已激活(累计上线 ≥ 1000): {activated_n} 家')


def main():
    conn = sqlite3.connect(DB)
    conn.executescript(DDL)
    load_gaode(conn, '/tmp/潜客明细表_20260516.xlsx', '2026-05-16')
    load_np(conn,    '/tmp/NP流转客户名单.xlsx',       '2026-05-24')
    verify(conn)
    conn.close()


if __name__ == '__main__':
    main()
