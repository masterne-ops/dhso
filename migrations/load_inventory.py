#!/usr/bin/env python3
"""代理商盘库明细入库 ``inventory_snapshot``。

兼容两类源文件：
1. 历史盘库明细（标题行 + 42列）；
2. Q3 起的紧凑盘库导出（表头在第1行，28列）。

安全策略：按盘库季度替换，不再全表 DELETE。历史季度保持不变；同一季度重复导入时，
按 ``(产品序列号, 盘库季度)`` 去重并取盘库时间最新一条。
"""
from __future__ import annotations

import argparse
import re
import sqlite3
from datetime import date, datetime
from pathlib import Path

import openpyxl


DEFAULT_DB = '/opt/so-data-analytics/db/product_flow.db'
DEFAULT_FILE = '/tmp/盘库明细.xlsx'

DB_COLUMNS = [
    '产品序列号', '盘库季度', '盘库活动名称', '盘库时间',
    '盘库客户省份', '盘库客户名称', '盘库客户编码',
    '上线时间', '上线自客户名称', '上线自客户编码', '安装红包时间',
    '产品名称', '外部型号', '内部型号', '产品料号',
    '产品一级', '产品二级', '研发一级', '产品三级', '产品四级',
    '产品系列', '产品子系列', '销售类型一级', '分销产品标签',
    '出货客户编码', '出货客户名称', '出库客户类型', '出库客户渠道客户类型',
    '出货客户省份', '出货客户城市', '出货客户区县', '出库时间',
    '合同类型', '分销合同类型', '合同编号', '订单编号', '订单时间',
    '产品下单价', '产品分销价', '产品现有分销价',
    '是否在库', '库龄天数', '导入时间',
]

PRICE_COLUMNS = {'产品下单价', '产品分销价', '产品现有分销价'}

# DB列名 -> 源文件可接受表头。多数同名，仅列出存在别名的字段。
SOURCE_ALIASES = {
    '产品序列号': ('产品序列号', '序列号'),
    '产品料号': ('产品料号', '物料号'),
    '出库时间': ('出库时间', '出货时间'),
}


def normalize_header(value) -> str:
    return re.sub(r'\s+', '', str(value or '')).strip()


def extract_quarter(activity) -> str:
    text = str(activity or '')
    match = re.search(r'(20\d{2})\s*Q\s*([1-4])', text, re.I)
    if match:
        return f'{match.group(1)}Q{match.group(2)}'
    if '25年' in text or '2025' in text:
        return '2025'
    return '其他'


def to_str(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    text = str(value).strip()
    return text or None


def to_num(value):
    if value is None or value == '':
        return None
    try:
        return float(str(value).replace(',', '').strip())
    except (ValueError, TypeError):
        return None


def parse_datetime(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value).strip()
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%Y/%m/%d %H:%M:%S', '%Y/%m/%d'):
        try:
            return datetime.strptime(text[:19], fmt)
        except ValueError:
            continue
    return None


def days_between(later, earlier):
    left, right = parse_datetime(later), parse_datetime(earlier)
    return (left - right).days if left and right else None


def find_header_and_iterator(ws):
    iterator = ws.iter_rows(values_only=True)
    for row_number in range(1, 11):
        try:
            row = next(iterator)
        except StopIteration:
            break
        headers = [normalize_header(v) for v in row]
        if '产品序列号' in headers or '序列号' in headers:
            return row_number, headers, iterator
    raise ValueError('前10行未找到“产品序列号/序列号”表头')


def build_source_index(headers):
    lookup = {name: idx for idx, name in enumerate(headers) if name}
    result = {}
    for column in DB_COLUMNS:
        if column in {'盘库季度', '是否在库', '库龄天数', '导入时间'}:
            continue
        candidates = SOURCE_ALIASES.get(column, (column,))
        result[column] = next((lookup[normalize_header(x)] for x in candidates
                               if normalize_header(x) in lookup), None)
    required = ['产品序列号', '盘库活动名称', '盘库时间', '盘库客户名称',
                '上线时间', '内部型号', '产品料号', '出库时间', '产品现有分销价']
    missing = [name for name in required if result.get(name) is None]
    if missing:
        raise ValueError(f'缺少必要字段：{missing}')
    return result


def get_value(row, index):
    return row[index] if index is not None and index < len(row) else None


def ensure_table(cursor):
    definitions = []
    for column in DB_COLUMNS:
        if column == '产品序列号':
            definitions.append('"产品序列号" TEXT NOT NULL')
        elif column == '盘库季度':
            definitions.append('"盘库季度" TEXT NOT NULL')
        elif column in PRICE_COLUMNS:
            definitions.append(f'"{column}" REAL')
        elif column in {'是否在库', '库龄天数'}:
            definitions.append(f'"{column}" INTEGER')
        else:
            definitions.append(f'"{column}" TEXT')
    cursor.execute(f"""
        CREATE TABLE IF NOT EXISTS inventory_snapshot (
            {','.join(definitions)},
            PRIMARY KEY (产品序列号, 盘库季度)
        )
    """)
    for name, column in (
        ('idx_inv_q', '盘库季度'), ('idx_inv_dealer', '盘库客户名称'),
        ('idx_inv_instock', '是否在库'), ('idx_inv_material', '产品料号'),
        ('idx_inv_series', '产品系列'), ('idx_inv_sn', '产品序列号'),
    ):
        cursor.execute(f'CREATE INDEX IF NOT EXISTS {name} ON inventory_snapshot("{column}")')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--db', default=DEFAULT_DB)
    parser.add_argument('--file', default=DEFAULT_FILE)
    parser.add_argument('--sheet', default=None)
    parser.add_argument('--quarter', default=None,
                        help='强制盘库季度，如2026Q3；会校验活动名称识别结果')
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()

    source = Path(args.file)
    print(f'读取 {source} ...')
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = workbook[args.sheet] if args.sheet else workbook[workbook.sheetnames[0]]
    header_row, headers, iterator = find_header_and_iterator(ws)
    source_index = build_source_index(headers)
    print(f'工作表 {ws.title}，表头第 {header_row} 行，识别 {sum(v is not None for v in source_index.values())} 个字段')

    dedup = {}
    raw_rows = blank_serials = 0
    detected_quarters = set()
    for row in iterator:
        raw_rows += 1
        serial = to_str(get_value(row, source_index['产品序列号']))
        if not serial:
            blank_serials += 1
            continue
        activity = get_value(row, source_index['盘库活动名称'])
        detected = extract_quarter(activity)
        detected_quarters.add(detected)
        if args.quarter:
            if detected not in {args.quarter, '其他'}:
                raise ValueError(f'活动名称识别为 {detected}，与 --quarter {args.quarter} 冲突')
            quarter = args.quarter
        else:
            quarter = detected
        inventory_time = to_str(get_value(row, source_index['盘库时间'])) or ''
        key = (serial, quarter)
        if key in dedup and inventory_time <= dedup[key][0]:
            continue
        dedup[key] = (inventory_time, row)
        if raw_rows % 100000 == 0:
            print(f'  已读 {raw_rows:,} 行，去重后 {len(dedup):,}')

    quarters = sorted({key[1] for key in dedup})
    print(f'总行 {raw_rows:,}，空序列号 {blank_serials:,}，去重后 {len(dedup):,}，季度 {quarters}')
    if not dedup:
        raise ValueError('没有可导入数据')
    if '其他' in quarters:
        raise ValueError('存在无法识别季度的数据，请用 --quarter 指定并核对活动名称')
    if args.dry_run:
        print('DRY RUN：未写入数据库')
        return

    imported_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    insert_columns = DB_COLUMNS
    placeholders = ','.join('?' for _ in insert_columns)
    column_sql = ','.join(f'"{name}"' for name in insert_columns)
    insert_sql = f'INSERT INTO inventory_snapshot ({column_sql}) VALUES ({placeholders})'

    connection = sqlite3.connect(args.db)
    cursor = connection.cursor()
    ensure_table(cursor)
    # 表/索引初始化可能开启隐式事务；先提交，再开启盘库季度替换事务。
    connection.commit()
    before_total = cursor.execute('SELECT COUNT(*) FROM inventory_snapshot').fetchone()[0]
    before_by_quarter = {
        q: cursor.execute('SELECT COUNT(*) FROM inventory_snapshot WHERE 盘库季度=?', (q,)).fetchone()[0]
        for q in quarters
    }

    try:
        cursor.execute('BEGIN IMMEDIATE')
        for quarter in quarters:
            cursor.execute('DELETE FROM inventory_snapshot WHERE 盘库季度=?', (quarter,))

        batch = []
        for (serial, quarter), (_, row) in dedup.items():
            values = {name: None for name in DB_COLUMNS}
            values['产品序列号'] = serial
            values['盘库季度'] = quarter
            for name, index in source_index.items():
                if index is None or name == '产品序列号':
                    continue
                raw = get_value(row, index)
                values[name] = to_num(raw) if name in PRICE_COLUMNS else to_str(raw)
            online_time = get_value(row, source_index['上线时间'])
            values['是否在库'] = 1 if not to_str(online_time) else 0
            values['库龄天数'] = days_between(
                get_value(row, source_index['盘库时间']),
                get_value(row, source_index['出库时间']),
            )
            values['导入时间'] = imported_at
            batch.append(tuple(values[name] for name in insert_columns))
            if len(batch) >= 5000:
                cursor.executemany(insert_sql, batch)
                batch.clear()
        if batch:
            cursor.executemany(insert_sql, batch)

        for quarter in quarters:
            inserted = cursor.execute(
                'SELECT COUNT(*) FROM inventory_snapshot WHERE 盘库季度=?', (quarter,)
            ).fetchone()[0]
            expected = sum(1 for key in dedup if key[1] == quarter)
            if inserted != expected:
                raise RuntimeError(f'{quarter} 写入行数 {inserted} != 预期 {expected}')
        connection.commit()
    except Exception:
        connection.rollback()
        raise

    after_total = cursor.execute('SELECT COUNT(*) FROM inventory_snapshot').fetchone()[0]
    print(f'\n✅ 数据库总行 {before_total:,} → {after_total:,}')
    for quarter in quarters:
        row = cursor.execute("""
            SELECT COUNT(*), SUM(是否在库),
                   ROUND(SUM(CASE WHEN 是否在库=1 THEN 产品现有分销价 ELSE 0 END)/10000.0, 1),
                   COUNT(DISTINCT 盘库客户名称), MIN(盘库时间), MAX(盘库时间)
              FROM inventory_snapshot WHERE 盘库季度=?
        """, (quarter,)).fetchone()
        print(f'{quarter}: 原 {before_by_quarter[quarter]:,} → {row[0]:,} 行；'
              f'在库 {int(row[1] or 0):,} 台；货值 {float(row[2] or 0):,.1f} 万；'
              f'代理商 {row[3]} 家；盘库时间 {row[4]} ~ {row[5]}')
    connection.close()


if __name__ == '__main__':
    main()
