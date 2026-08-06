#!/usr/bin/env python3
"""product_flow 全量替换 — 用一个或多个 FX601 分片替换指定日期/月度范围。

根治 FX601 重复:该表 ID 不稳定(同一次激活在多批导出里 ID 不同、连格式都变)、
产品序列号有 *** 占位(外省代理商在本省出货),没有可靠去重键。按 ID 的
INSERT OR IGNORE 拦不住"同激活不同ID",导致多批导出累积重复(SO 虚高)。

全量替换不依赖去重键:单次 FX601 导出 = 该时点真实快照(每行一台真实设备、
*** 各自独立行、无重复),DELETE 覆盖月份 + INSERT 全量,天然规避累积。

用法:
  python load_pf_replace.py --file /tmp/FX601_1.xlsx --file /tmp/FX601_2.xlsx \
      --dates 2025-01-01:2026-07-25
  python load_pf_replace.py --file /tmp/FX601_1.xlsx --file /tmp/FX601_2.xlsx \
      --dates 2025-01-01:2026-07-25 --apply
"""
import sqlite3, math, argparse
from datetime import date, datetime
from openpyxl import load_workbook

DB = '/opt/so-data-analytics/db/product_flow.db'


def to_cell(v):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    return v


def is_blank(c):
    return c is None or (isinstance(c, float) and math.isnan(c))


def find_header(rows, must_have, scan=8):
    for i, row in enumerate(rows[:scan]):
        cells = set(str(c).strip() for c in row if not is_blank(c))
        if all(k in cells for k in must_have):
            return i
    raise ValueError(f"找不到表头(需含 {must_have})")


def date_text(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ''
    return str(v).strip()[:10]


def open_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    head_rows = []
    for _ in range(8):
        try:
            head_rows.append(next(rows))
        except StopIteration:
            break
    h = find_header(head_rows, ['ID', '产品序列号', '上线时间'])
    header = [str(c).strip() if not is_blank(c) else None for c in head_rows[h]]

    def iterator():
        try:
            for row in head_rows[h + 1:]:
                yield row
            yield from rows
        finally:
            wb.close()

    return ws.title, h, header, iterator()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', action='append', required=True,
                    help='FX601 全量导出 xlsx；分片文件可重复传入')
    group = ap.add_mutually_exclusive_group(required=True)
    group.add_argument('--dates', help='替换的上线日期范围 LO:HI，如 2025-01-01:2026-07-25')
    group.add_argument('--months', help='替换的上线年月范围 LO:HI，如 2026-01:2026-05')
    ap.add_argument('--db', default=DB, help=f'SQLite 数据库路径（默认 {DB}）')
    ap.add_argument('--apply', action='store_true', help='正式执行;不加则 dry-run')
    a = ap.parse_args()
    LO, HI = (a.dates or a.months).split(':')
    exact_dates = bool(a.dates)
    conn = sqlite3.connect(a.db)
    dbc = [r[1] for r in conn.execute('PRAGMA table_info(product_flow)')]
    period_expr = "date(上线时间)" if exact_dates else "substr(上线时间,1,7)"
    before = conn.execute(f"SELECT COUNT(*) FROM product_flow WHERE {period_expr} BETWEEN ? AND ?",
                          (LO, HI)).fetchone()[0]

    def iter_values():
        for path in a.file:
            sheet, h, header, rows = open_rows(path)
            cm = [(i, c) for i, c in enumerate(header) if c and c in dbc]
            if set(dbc) - {c for _, c in cm}:
                missing = sorted(set(dbc) - {c for _, c in cm})
                raise ValueError(f"{path}: 缺少数据库字段 {missing}")
            ti = header.index('上线时间')
            print(f"{path} | sheet={sheet} | 表头第{h+1}行 | 列对齐 {len(cm)}/{len(dbc)}")
            for row in rows:
                if not any(not is_blank(c) for c in row):
                    continue
                dt = date_text(row[ti] if ti < len(row) else None)
                key = dt if exact_dates else dt[:7]
                if LO <= key <= HI:
                    yield cm, tuple(to_cell(row[i]) if i < len(row) else None for i, _ in cm)

    source_count = 0
    out_of_range = 0
    if not a.apply:
        for _, _row in iter_values():
            source_count += 1
        print(f"源文件命中[{LO}~{HI}]: {source_count} 行")
        print(f"生产现有[{LO}~{HI}]: {before} 行 → 将删 {before} + 插 {source_count} "
              f"(净变化 {source_count - before:+d})")
        print("\n[DRY-RUN] 未写库。核对无误后加 --apply 执行(单事务,失败回滚)。")
        conn.close()
        return

    conn.execute('BEGIN IMMEDIATE')
    try:
        conn.execute(f"DELETE FROM product_flow WHERE {period_expr} BETWEEN ? AND ?", (LO, HI))
        batch, batch_cm = [], None
        for cm, row in iter_values():
            current_cm = tuple(c for _, c in cm)
            if batch_cm is None:
                batch_cm = current_cm
            if current_cm != batch_cm:
                raise ValueError("分片字段顺序不一致")
            batch.append(row)
            if len(batch) >= 5000:
                csql = ', '.join(f'"{c}"' for c in batch_cm)
                ph = ', '.join(['?'] * len(batch_cm))
                conn.executemany(f'INSERT INTO product_flow ({csql}) VALUES ({ph})', batch)
                source_count += len(batch)
                batch.clear()
        if batch:
            csql = ', '.join(f'"{c}"' for c in batch_cm)
            ph = ', '.join(['?'] * len(batch_cm))
            conn.executemany(f'INSERT INTO product_flow ({csql}) VALUES ({ph})', batch)
            source_count += len(batch)
        after = conn.execute(f"SELECT COUNT(*) FROM product_flow WHERE {period_expr} BETWEEN ? AND ?",
                             (LO, HI)).fetchone()[0]
        if after != source_count:
            raise RuntimeError(f"写入后行数不一致: source={source_count}, db={after}")
        conn.commit()
        print(f"✅ 替换完成,[{LO}~{HI}] {before} → {after} 行（净变化 {after-before:+d}）")
    except Exception as e:
        conn.rollback()
        print(f"❌ 出错,已整体回滚: {e}")
        raise
    finally:
        conn.close()


if __name__ == '__main__':
    main()
