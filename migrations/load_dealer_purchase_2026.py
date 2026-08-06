#!/usr/bin/env python3
"""代理商进货明细 2026 年度替换入库 → dealer_purchase(供 page40 代理商进货指导)。

源: 「工作簿3.xlsx」类导出(单 Sheet,14 列,无年份/区县/计销售列,日期为「N月D日」文本):
  日期 / 业绩归属二级办 / 业务员姓名 / 下单客户 / 行业一级 / 产品一~四级 /
  内部型号 / 外部型号 / 求和项:实销（万) / 求和项:实发数量 / 求和项:计奖励

策略: **2026 整年替换** —— DELETE 数据年份=2026 + INSERT 文件全量(2025 不动)。
天然覆盖"之前已入过的 2026 部分月份"(旧 26 年 1-3 月数据被本文件更完整版本替换)。
下单客户类型派生与 load_dealer_purchase.py 同口径:
  一级代理商=命中 dealer_si_snapshot(+SMB 补充名单) / 服务商=命中 provider_contract / 其他。

用法:
  python load_dealer_purchase_2026.py --file /tmp/工作簿3.xlsx           # dry-run
  python load_dealer_purchase_2026.py --file /tmp/工作簿3.xlsx --apply   # 正式(单事务)
"""
import sqlite3
import argparse
import openpyxl
from datetime import datetime

DB = '/opt/so-data-analytics/db/product_flow.db'
YEAR = 2026

# 与 load_dealer_purchase.py 保持一致的 SMB 补充名单(人工确认属一级代理商)
SMB_SUPPLEMENT = [
    '浙江名硕建设工程有限公司',
    '浙江歌呈智能科技有限公司',
    '微空间网络科技（湖州）有限公司',
    '湖州微空间网络科技有限公司',
]

# Excel 列名 → 位置在读表头时确定;这里定义 Excel 列名 → 表列名
COLMAP = {
    '日期': '日期原文',
    '业绩归属二级办': '业绩归属二级办',
    '业务员姓名': '业务员姓名',
    '下单客户': '下单客户',
    '行业一级': '行业一级',
    '产品一级': '产品一级', '产品二级': '产品二级',
    '产品三级': '产品三级', '产品四级': '产品四级',
    '内部型号': '内部型号', '外部型号': '外部型号',
    '求和项:实销（万)': '实销万',
    '求和项:实发数量': '实发数量',
    '求和项:计奖励': '计奖励',
}


def month_of(d):
    s = str(d or '')
    if '月' in s:
        try:
            return int(s.split('月')[0])
        except ValueError:
            return None
    return None


def num(v):
    if v is None or v == '':
        return None
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', required=True)
    ap.add_argument('--apply', action='store_true')
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.file, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else None for c in rows[0]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    miss = [e for e in COLMAP if e not in idx]
    if miss:
        raise SystemExit(f'❌ 文件缺列: {miss}')

    batch, by_month = [], {}
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    for r in rows[1:]:
        if not r or not r[idx['下单客户']]:
            continue
        m = month_of(r[idx['日期']])
        by_month[m] = by_month.get(m, 0) + 1
        batch.append((
            YEAR, m, str(r[idx['日期']] or '').strip(),
            str(r[idx['业绩归属二级办']] or '').strip() or None,
            str(r[idx['业务员姓名']] or '').strip() or None,
            str(r[idx['下单客户']] or '').strip(),
            None,  # 下单客户类型,入库后 UPDATE 派生
            None,  # 下单客户区县(本文件无)
            str(r[idx['行业一级']] or '').strip() or None,
            str(r[idx['产品一级']] or '').strip() or None,
            str(r[idx['产品二级']] or '').strip() or None,
            str(r[idx['产品三级']] or '').strip() or None,
            str(r[idx['产品四级']] or '').strip() or None,
            str(r[idx['内部型号']] or '').strip() or None,
            str(r[idx['外部型号']] or '').strip() or None,
            num(r[idx['求和项:实销（万)']]),
            num(r[idx['求和项:实发数量']]),
            None,  # 计返利(本文件无)
            None,  # 计销售(本文件无;已确认代码无消费方)
            num(r[idx['求和项:计奖励']]),
            now,
        ))

    conn = sqlite3.connect(DB)
    old = conn.execute('SELECT COUNT(*) FROM dealer_purchase WHERE 数据年份=?', (YEAR,)).fetchone()[0]
    amt = sum(v[15] or 0 for v in batch)
    print(f'文件: {len(batch):,} 行 / 实销 {amt:,.1f} 万 / 月份分布 {dict(sorted(by_month.items()))}')
    print(f'库中 {YEAR} 年现有 {old:,} 行 → 将整年替换')
    if not a.apply:
        print('[DRY-RUN] 加 --apply 执行(单事务,失败回滚)。')
        conn.close()
        return

    cols = ['数据年份', '月份', '日期原文', '业绩归属二级办', '业务员姓名', '下单客户', '下单客户类型',
            '下单客户区县', '行业一级', '产品一级', '产品二级', '产品三级', '产品四级',
            '内部型号', '外部型号', '实销万', '实发数量', '计返利', '计销售', '计奖励', '导入时间']
    ph = ','.join('?' * len(cols))
    conn.execute('BEGIN')
    try:
        conn.execute('DELETE FROM dealer_purchase WHERE 数据年份=?', (YEAR,))
        conn.executemany(f"INSERT INTO dealer_purchase ({','.join(cols)}) VALUES ({ph})", batch)
        # 派生下单客户类型(只对本年,口径同 load_dealer_purchase)
        conn.execute("""
            UPDATE dealer_purchase SET 下单客户类型='一级代理商'
             WHERE 数据年份=? AND 下单客户 IN (SELECT DISTINCT "客户名称（无@办事处）" FROM dealer_si_snapshot)
        """, (YEAR,))
        ph2 = ','.join('?' * len(SMB_SUPPLEMENT))
        conn.execute(f"UPDATE dealer_purchase SET 下单客户类型='一级代理商' "
                     f"WHERE 数据年份=? AND 下单客户 IN ({ph2})", (YEAR, *SMB_SUPPLEMENT))
        conn.execute("""
            UPDATE dealer_purchase SET 下单客户类型='服务商'
             WHERE 数据年份=? AND 下单客户类型 IS NULL
               AND 下单客户 IN (SELECT DISTINCT 客户名称 FROM provider_contract)
        """, (YEAR,))
        conn.execute("UPDATE dealer_purchase SET 下单客户类型='其他' "
                     "WHERE 数据年份=? AND 下单客户类型 IS NULL", (YEAR,))
        conn.commit()
    except Exception as e:
        conn.rollback()
        print(f'❌ 出错,已整体回滚: {e}')
        raise

    print(f'✅ 替换完成: 删旧 {old:,} + 插 {len(batch):,}')
    for r in conn.execute("""
        SELECT 月份, 下单客户类型, COUNT(*), ROUND(SUM(实销万),1)
          FROM dealer_purchase WHERE 数据年份=? GROUP BY 1,2 ORDER BY 1,4 DESC
    """, (YEAR,)):
        print(f'   {r[0]}月 {r[1]:6s} {r[2]:>6,} 行 / {r[3]:>8,.1f} 万')
    conn.close()


if __name__ == '__main__':
    main()
