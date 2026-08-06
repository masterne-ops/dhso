#!/usr/bin/env python3
"""NP 转入客户 np_transfer_customer 入库 + 合并旧 NP 流转名单。

背景:
  - NP 转入客户 = 客户来源「NP流转/一站式」的潜客(从公共池转入我方跟进)。
  - 源文件「潜客明细表_YYYYMMDD.xlsx」(46 列全字段),表头在第 2 行(行0是标题)。
  - 现有旧 NP 流转名单在 np_customer_pool(6 列简版),按公共字段合并进来。

统一表 np_transfer_customer:46 列全字段 + 数据时点 + 来源批次(客户来源),
按数据时点快照(保留历史、看进展)。管理者看板 page 读它。

进展漏斗:转入 → 已报备 → 已转出 → 已签约(客户名称命中 provider_contract) → 已激活(累计上线≥1000)。

用法:
  python load_np_transfer.py --file /tmp/潜客明细表_20260705.xlsx --period 2026-07-05           # dry-run
  python load_np_transfer.py --file /tmp/潜客明细表_20260705.xlsx --period 2026-07-05 --apply     # 入新文件
  python load_np_transfer.py --merge-old --apply                                                 # 另把旧 np_customer_pool 合并进来
"""
import sqlite3
import argparse
import openpyxl
from datetime import datetime

DB = '/opt/so-data-analytics/db/product_flow.db'

# np_transfer_customer 字段(与潜客明细表一致,46 列 + 数据时点 + 来源批次)
COLS = [
    '外部客户名称', '客户来源', '客户编码', '客户名称', '是否已转出', '与我司业务相关',
    '下一步计划', '原因说明', '预计完成时间', '是否转出组织', '签约类型', '客户状态',
    '客户所有者工号', '客户所有者', '客户所有者部门', '年安防采购量', '对大华品牌认可度',
    '责任分销商', '责任分销商分配人', '责任分销商分配时间', '责任分销经理', '责任分销经理工号',
    '责任分销经理分配人', '责任分销经理分配时间', '客户联系人电话', '客户联系地址',
    '省', '市', '区县', '分销商认证', '渠道客户类型', '协同人姓名', '安装红包上线金额',
    '最近拜访时间_分销经理', '年度拜访次数_分销经理', '本月拜访次数_分销经理',
    '最近拜访时间_分销商', '年度拜访次数_分销商', '本月拜访次数_分销商',
    '创建时间', '外部客户关联时间', '客户关闭时间', '转公共池时间', '签约时间',
    '首次梳理时间', '最后1次梳理时间',
]
_REAL = {'安装红包上线金额'}
_INT = {'年度拜访次数_分销经理', '本月拜访次数_分销经理', '年度拜访次数_分销商', '本月拜访次数_分销商'}

# 潜客 Excel 列名(带换行) → DB 列名
RENAME = {
    '最近拜访时间\n（分销经理）': '最近拜访时间_分销经理',
    '年度拜访次数\n（分销经理）': '年度拜访次数_分销经理',
    '本月拜访次数\n（分销经理）': '本月拜访次数_分销经理',
    '最近拜访时间\n（分销商）': '最近拜访时间_分销商',
    '年度拜访次数\n（分销商）': '年度拜访次数_分销商',
    '本月拜访次数\n（分销商）': '本月拜访次数_分销商',
}


def _ddl():
    defs = ['数据时点 TEXT NOT NULL']
    for c in COLS:
        ty = 'REAL' if c in _REAL else ('INTEGER' if c in _INT else 'TEXT')
        defs.append(f'"{c}" {ty}')
    return f"""
    CREATE TABLE IF NOT EXISTS np_transfer_customer (
        {','.join(defs)},
        PRIMARY KEY (数据时点, 外部客户名称)
    );
    CREATE INDEX IF NOT EXISTS idx_npt_city ON np_transfer_customer(市);
    CREATE INDEX IF NOT EXISTS idx_npt_source ON np_transfer_customer(客户来源);
    CREATE INDEX IF NOT EXISTS idx_npt_signed ON np_transfer_customer(客户名称);
    CREATE INDEX IF NOT EXISTS idx_npt_mgr ON np_transfer_customer(责任分销经理);
    """


def to_cell(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, str):
        return v.strip()
    return v


def load_new_file(conn, fp, period, apply):
    print(f'\n══ NP 转入(潜客明细) {fp} → 数据时点 {period} ══')
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    # 表头行:找含「外部客户名称」的行
    hr = next(i for i, r in enumerate(rows[:5])
              if r and '外部客户名称' in [str(x).strip() if x else '' for x in r])
    hdr = [RENAME.get(str(c).strip(), str(c).strip()) if c else None for c in rows[hr]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    use = [c for c in COLS if c in idx]
    miss = [c for c in COLS if c not in idx]
    print(f'  表头第{hr+1}行 | 数据 {len(rows)-hr-1} 行 | 对齐 {len(use)}/{len(COLS)} 列')
    if miss:
        print(f'  ⚠️ 缺列(留空): {miss}')

    batch = []
    for r in rows[hr + 1:]:
        if not r or not any(r):
            continue
        ext = r[idx['外部客户名称']] if idx.get('外部客户名称') is not None and idx['外部客户名称'] < len(r) else None
        if not ext:
            continue
        vals = [period] + [to_cell(r[idx[c]]) if idx[c] < len(r) else None for c in use]
        batch.append(tuple(vals))
    print(f'  有效行(外部客户名称非空): {len(batch)}')
    if not apply:
        print('  [DRY-RUN] 加 --apply 写库。')
        return
    csql = ', '.join(f'"{c}"' for c in ['数据时点'] + use)
    ph = ', '.join(['?'] * (1 + len(use)))
    conn.execute('DELETE FROM np_transfer_customer WHERE 数据时点=?', (period,))
    conn.executemany(f'INSERT OR IGNORE INTO np_transfer_customer ({csql}) VALUES ({ph})', batch)
    conn.commit()
    n = conn.execute('SELECT COUNT(*) FROM np_transfer_customer WHERE 数据时点=?', (period,)).fetchone()[0]
    print(f'  ✅ 入库 {n} 行')


def merge_old(conn, apply):
    """把旧 np_customer_pool(简版) 按公共字段映射进 np_transfer_customer。
    映射:客户名称(外部公司名)→外部客户名称;签约服务商名称→客户名称;地市→市;区县→区县;
         区县所有者负责人→责任分销经理;客户来源填 'NP流转名单'。数据时点沿用 np_customer_pool 的时点。
    """
    src = conn.execute("""
        SELECT 数据时点, 客户名称, 地市, 区县, 区县所有者负责人, 签约服务商名称
          FROM np_customer_pool
    """).fetchall()
    print(f'\n══ 合并旧 NP 流转名单 np_customer_pool: {len(src)} 行 ══')
    if not apply:
        print('  [DRY-RUN] 加 --apply 写库。')
        return
    cols = ['数据时点', '外部客户名称', '客户名称', '市', '区县', '责任分销经理', '客户来源']
    csql = ', '.join(f'"{c}"' for c in cols)
    ph = ', '.join(['?'] * len(cols))
    batch = [(dt, name, signed, city, dist, owner, 'NP流转名单')
             for dt, name, city, dist, owner, signed in src if name]
    conn.executemany(f'INSERT OR IGNORE INTO np_transfer_customer ({csql}) VALUES ({ph})', batch)
    conn.commit()
    n = conn.execute("SELECT COUNT(*) FROM np_transfer_customer WHERE 客户来源='NP流转名单'").fetchone()[0]
    print(f'  ✅ 合并 {n} 行(客户来源=NP流转名单)')


def verify(conn):
    print('\n═══ NP 转入进展(全表) ═══')
    total = conn.execute('SELECT COUNT(*) FROM np_transfer_customer').fetchone()[0]
    reported = conn.execute("SELECT COUNT(*) FROM np_transfer_customer WHERE 客户状态='已报备'").fetchone()[0]
    transferred = conn.execute("SELECT COUNT(*) FROM np_transfer_customer WHERE 是否已转出='Y'").fetchone()[0]
    signed = conn.execute("""
        SELECT COUNT(DISTINCT t.外部客户名称) FROM np_transfer_customer t
         WHERE t.客户名称 IS NOT NULL AND t.客户名称!=''
           AND EXISTS (SELECT 1 FROM provider_contract pc WHERE pc.客户名称=t.客户名称)
    """).fetchone()[0]
    activated = conn.execute("""
        WITH sg AS (SELECT DISTINCT 客户名称 FROM np_transfer_customer WHERE 客户名称 IS NOT NULL AND 客户名称!=''),
             amt AS (SELECT 上线客户名称, SUM(产品现有分销价) v FROM install_redpack GROUP BY 1)
        SELECT COUNT(*) FROM sg JOIN amt a ON a.上线客户名称=sg.客户名称 WHERE a.v>=1000
    """).fetchone()[0]
    print(f'  转入 {total} → 已报备 {reported} → 已转出 {transferred} → 已签约 {signed} → 已激活 {activated}')
    print('  按来源批次:')
    for r in conn.execute("SELECT 客户来源, 数据时点, COUNT(*) FROM np_transfer_customer GROUP BY 1,2 ORDER BY 2 DESC"):
        print(f'    {r[0]} ({r[1]}): {r[2]} 家')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file')
    ap.add_argument('--period')
    ap.add_argument('--merge-old', action='store_true')
    ap.add_argument('--apply', action='store_true')
    a = ap.parse_args()
    conn = sqlite3.connect(DB)
    conn.executescript(_ddl())
    if a.file:
        assert a.period, '--file 需配 --period'
        load_new_file(conn, a.file, a.period, a.apply)
    if a.merge_old:
        merge_old(conn, a.apply)
    if a.apply:
        verify(conn)
    conn.close()


if __name__ == '__main__':
    main()
