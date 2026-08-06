#!/usr/bin/env python3
"""智能备货配单：自动识别任务进度 → 型号/价格/整数数量。"""
from __future__ import annotations

import io
import sys
from copy import copy
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent.parent))
from _auth import filter_by_scope, require_auth  # noqa: E402
from _loaders import DB_PATH  # noqa: E402
from _si_order_tool import (  # noqa: E402
    allocation_status,
    auto_allocate,
    enrich_order_rows,
    load_tool_dataset,
    reallocate_existing,
    validate_dataset,
)


require_auth()

st.markdown(
    """
    <style>
    .block-container {padding-top: 1.4rem; padding-bottom: 3rem;}
    [data-testid="stMetricValue"] {font-size: 1.65rem;}
    .si-note {padding:.8rem 1rem;border-radius:10px;background:#f3f7fd;
              border:1px solid #d9e5f6;color:#24476f;margin:.35rem 0 1rem;}
    .si-good {color:#0d6948;font-weight:700;}
    .si-warn {color:#a15c00;font-weight:700;}
    </style>
    """,
    unsafe_allow_html=True,
)


@st.cache_data(ttl=600, show_spinner="读取最新任务、SI、SO和分销报价…")
def _load_dataset(target_month: int | None = None):
    return load_tool_dataset(DB_PATH, target_month=target_month)


def _state_key(year: int, target_month: int, code: str) -> str:
    return f"si_order_tool_rows_{year}_{target_month}_{code}"


def _as_state_records(df: pd.DataFrame) -> list[dict]:
    if df.empty:
        return []
    return df[["内部型号", "建议下单数量"]].to_dict("records")


def _export_xlsx(customer: pd.Series, order: pd.DataFrame, status: dict, meta: dict) -> bytes:
    target_month = int(meta["target_month"])
    summary = pd.DataFrame(
        [
            ["客户名称", customer["客户名称"]],
            ["客户编码", customer["客户编码"]],
            ["城市", customer["城市"]],
            ["负责人", customer["客户所有者"]],
            ["客户类型", customer["客户类型"]],
            ["年度任务（万）", customer["年度任务_万"]],
            ["累计业绩达成（计任务，万）", customer["当前官方达成_万"]],
            [f"{target_month}月底累计任务（万）", customer["目标月底累计任务_万"]],
            [f"距{target_month}月底目标完成度", customer["当前进度"]],
            [f"{target_month}月下单缺口（万）", status["target_wan"]],
            ["建议下单金额（万）", status["amount_wan"]],
            ["配单差异（万）", status["difference_wan"]],
            ["SI数据月份", meta["si_latest_month"]],
            ["SO截止日", meta["so_asof"]],
            ["报价日期", meta["quotation_date"]],
        ],
        columns=["项目", "数值"],
    )
    detail_cols = [
        "标准产品四级",
        "内部型号",
        "外部型号",
        "SI_YTD数量",
        "SI_近3月数量",
        "SO_YTD数量",
        "SO_90数量",
        "近期月均SO",
        "最新分销价_元",
        "建议下单数量",
        "建议下单金额_万",
        "价格来源",
        "产品状态",
        "建议依据",
        "最近SO日期",
    ]
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="任务概览", index=False)
        order[[c for c in detail_cols if c in order.columns]].to_excel(
            writer, sheet_name="建议下单明细", index=False
        )
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(
                    name=copy(cell.font).name,
                    size=copy(cell.font).size,
                    bold=True,
                    color="FFFFFF",
                )
                cell.fill = PatternFill(fill_type="solid", fgColor="2F75B5")
    return output.getvalue()


st.markdown("## 🧮 代理商智能备货配单")
st.markdown(
    '<div class="si-note">系统自动识别正式任务当前做到的月份，并默认规划下一备货周期。'
    "选择客户 → 查看任务缺口 → 智能配单；"
    "表中可直接新增、删除或修改型号与数量，型号变化后自动带出最新可验证价格。"
    "所有编辑仅保存在当前会话，不回写生产业务库。</div>",
    unsafe_allow_html=True,
)

if st.button("🔄 刷新生产数据", help="清除10分钟缓存并重新读取生产库"):
    _load_dataset.clear()
    st.rerun()

base_dataset = _load_dataset()
base_meta = base_dataset["meta"]
target_options = [int(x) for x in base_meta["available_target_months"]]
target_month = st.selectbox(
    "备货周期",
    target_options,
    index=target_options.index(int(base_meta["default_target_month"])),
    format_func=lambda value: f'{base_meta["year"]}年{value}月',
    help=(
        f'系统根据正式累计任务识别当前任务已做到{base_meta["task_through_month"]}月，'
        "默认选择下一期；可切换到后续月份做预案。"
    ),
)
dataset = (
    base_dataset
    if int(target_month) == int(base_meta["target_month"])
    else _load_dataset(int(target_month))
)
meta = dataset["meta"]
target_year = int(meta["year"])
target_month = int(meta["target_month"])
checks = validate_dataset(dataset)
failed = [x for x in checks if not x["通过"]]
if failed:
    st.error("数据校验未通过：" + "；".join(f'{x["检查"]}（{x["说明"]}）' for x in failed))
    st.stop()

customers = filter_by_scope(
    dataset["customers"],
    city_col="城市",
    district_col="区县",
    salesperson_col="客户所有者",
    dealer_col="客户名称",
).copy()
if customers.empty:
    st.warning("当前账号的数据权限范围内没有签约客户。")
    st.stop()

filters = st.columns([1.2, 1.2, 3.2])
city_options = ["全部"] + sorted(customers["城市"].dropna().astype(str).unique())
city = filters[0].selectbox("城市", city_options)
city_view = customers if city == "全部" else customers[customers["城市"].eq(city)]
owner_options = ["全部"] + sorted(city_view["客户所有者"].dropna().astype(str).unique())
owner = filters[1].selectbox("负责人", owner_options)
view = city_view if owner == "全部" else city_view[city_view["客户所有者"].eq(owner)]
view = view.sort_values(["目标月下单缺口_万", "客户名称"], ascending=[False, True])
customer_map = view.set_index("客户编码").to_dict("index")
codes = view["客户编码"].astype(str).tolist()
code = filters[2].selectbox(
    "客户",
    codes,
    format_func=lambda x: (
        f'{customer_map[x]["客户名称"]}｜缺口 '
        f'{customer_map[x]["目标月下单缺口_万"]:,.2f} 万'
    ),
)
customer = view.loc[view["客户编码"].astype(str).eq(str(code))].iloc[0]
customer_models = dataset["models"][
    dataset["models"]["客户编码"].astype(str).eq(str(code))
].copy()
catalog = dataset["catalog"].copy()
target_wan = float(customer["目标月下单缺口_万"])

kpi = st.columns(5)
kpi[0].metric("年度任务", f'{customer["年度任务_万"]:,.0f} 万')
kpi[1].metric("累计业绩达成（计任务）", f'{customer["当前官方达成_万"]:,.2f} 万')
kpi[2].metric(f"{target_month}月底累计任务", f'{customer["目标月底累计任务_万"]:,.2f} 万')
kpi[3].metric(f"距{target_month}月底目标", f'{customer["当前进度"]:.1%}')
kpi[4].metric(f"{target_month}月下单缺口", f"{target_wan:,.2f} 万")
st.progress(
    min(max(float(customer["当前进度"]), 0), 1.0),
    text=f'{target_month}月底累计任务进度：{customer["当前进度"]:.1%}',
)
st.caption(
    f'负责人：{customer["客户所有者"]}　｜　客户类型：{customer["客户类型"]}'
    f'　｜　SI覆盖：{customer["SI覆盖"]}'
)
if customer["SI覆盖"] != "SI已匹配":
    st.warning("该客户在生产 SI 明细中没有有效型号记录，智能配单主要参考 SO；请补充核对 SI。")

state_key = _state_key(target_year, target_month, str(code))
if state_key not in st.session_state:
    initial = auto_allocate(customer_models, target_wan)
    st.session_state[state_key] = _as_state_records(initial)

st.markdown("### 建议下单型号")
action_cols = st.columns([1.1, 1.1, 1.1, 3.2, 1.0])
if action_cols[0].button("✨ 智能配单", type="primary", width="stretch"):
    suggested = auto_allocate(customer_models, target_wan)
    st.session_state[state_key] = _as_state_records(suggested)
    st.rerun()

current_order = enrich_order_rows(
    st.session_state[state_key], customer_models, catalog
)
if action_cols[1].button("🎯 当前型号对齐", width="stretch"):
    fitted = (
        reallocate_existing(current_order, target_wan)
        if not current_order.empty
        else auto_allocate(customer_models, target_wan)
    )
    st.session_state[state_key] = _as_state_records(fitted)
    st.rerun()
if action_cols[2].button("🧹 清空", width="stretch"):
    st.session_state[state_key] = []
    st.rerun()

catalog = catalog[catalog["内部型号"].fillna("").astype(str).ne("")].copy()
catalog_options = catalog["内部型号"].astype(str).tolist()
catalog_lookup = catalog.set_index("内部型号").to_dict("index")


def _model_label(model: str) -> str:
    row = catalog_lookup[model]
    raw_price = row.get("最新分销价_元")
    price_text = f'¥{float(raw_price):,.2f}' if pd.notna(raw_price) else "无可验证价格"
    return f'{model}｜{row.get("目录外部型号") or "-"}｜{price_text}'


selected_model = action_cols[3].selectbox(
    "新增型号（可输入搜索）",
    options=catalog_options,
    index=None,
    placeholder="输入内部型号搜索…",
    format_func=_model_label,
    label_visibility="collapsed",
)
if action_cols[4].button(
    "＋ 添加", width="stretch", disabled=not selected_model
):
    rows = list(st.session_state[state_key])
    if selected_model and not any(r.get("内部型号") == selected_model for r in rows):
        rows.append({"内部型号": selected_model, "建议下单数量": 0})
    st.session_state[state_key] = rows
    st.rerun()

current_order = enrich_order_rows(
    st.session_state[state_key], customer_models, catalog
)
display_cols = [
    "删除",
    "标准产品四级",
    "内部型号",
    "外部型号",
    "SI_YTD数量",
    "SO_90数量",
    "近期月均SO",
    "最新分销价_元",
    "建议下单数量",
    "建议下单金额_万",
    "价格来源",
    "产品状态",
    "建议依据",
]
editor_df = current_order.copy()
editor_df.insert(0, "删除", False)
for col in display_cols:
    if col not in editor_df.columns:
        editor_df[col] = None
editor_df = editor_df[display_cols]
edited = st.data_editor(
    editor_df,
    width="stretch",
    hide_index=True,
    height=min(640, 78 + max(len(editor_df), 1) * 36),
    num_rows="dynamic",
    disabled=[
        "标准产品四级",
        "外部型号",
        "SI_YTD数量",
        "SO_90数量",
        "近期月均SO",
        "最新分销价_元",
        "建议下单金额_万",
        "价格来源",
        "产品状态",
        "建议依据",
    ],
    column_config={
        "删除": st.column_config.CheckboxColumn("删除", width="small"),
        "标准产品四级": st.column_config.TextColumn(width="small"),
        "内部型号": st.column_config.SelectboxColumn(
            "内部型号", options=catalog_options, required=True, width="large"
        ),
        "外部型号": st.column_config.TextColumn(width="large"),
        "SI_YTD数量": st.column_config.NumberColumn("SI年内", format="%.0f"),
        "SO_90数量": st.column_config.NumberColumn("SO近90天", format="%.0f"),
        "近期月均SO": st.column_config.NumberColumn(format="%.1f"),
        "最新分销价_元": st.column_config.NumberColumn("分销价（元）", format="%.2f"),
        "建议下单数量": st.column_config.NumberColumn(
            "下单数量", min_value=0, step=1, format="%d", required=True
        ),
        "建议下单金额_万": st.column_config.NumberColumn("金额（万）", format="%.2f"),
        "价格来源": st.column_config.TextColumn(width="medium"),
        "产品状态": st.column_config.TextColumn(width="small"),
        "建议依据": st.column_config.TextColumn(width="large"),
    },
    key=f"si_order_editor_{code}",
)

edited_records = []
for row in pd.DataFrame(edited).to_dict("records"):
    if bool(row.get("删除")) or not str(row.get("内部型号") or "").strip():
        continue
    edited_records.append(
        {
            "内部型号": str(row["内部型号"]).strip(),
            "建议下单数量": row.get("建议下单数量") or 0,
        }
    )
new_order = enrich_order_rows(edited_records, customer_models, catalog)
old_pairs = _as_state_records(current_order)
new_pairs = _as_state_records(new_order)
if old_pairs != new_pairs:
    st.session_state[state_key] = new_pairs
    st.rerun()

status = allocation_status(current_order, target_wan)
totals = st.columns(4)
totals[0].metric("建议下单金额", f'{status["amount_wan"]:,.2f} 万')
totals[1].metric(
    "与缺口差异",
    f'{status["difference_wan"]:+,.2f} 万',
    delta=f'容差 {status["tolerance_wan"]:.2f} 万',
    delta_color="off",
)
totals[2].metric("建议台数", f'{status["unit_count"]:,} 台')
totals[3].metric("建议型号", f'{status["model_count"]} 个')
st.progress(
    min(max(status["coverage"], 0), 1.0),
    text=f'下单金额匹配度：{status["coverage"]:.1%}',
)
if target_wan <= 0:
    st.success(f"当前已覆盖{target_month}月底累计任务，不需要为任务追单。")
elif status["aligned"]:
    st.success(
        f'✅ 已基本对齐：建议 {status["amount_wan"]:,.2f} 万，'
        f'差异 {status["difference_wan"]:+,.2f} 万。'
    )
elif status["difference_wan"] < 0:
    st.warning(f'还差 {status["remaining_wan"]:,.2f} 万，可继续加型号或点击“当前型号对齐”。')
else:
    st.warning(f'当前超出缺口 {status["difference_wan"]:,.2f} 万，可下调数量后自动重算。')

download = _export_xlsx(customer, current_order, status, dataset["meta"])
st.download_button(
    "⬇️ 导出该客户下单建议 Excel",
    data=download,
    file_name=f'{customer["客户名称"]}-{target_year}年{target_month}月备货建议.xlsx',
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)

with st.expander("数据口径与待核实项", expanded=False):
    period_ratio = float(customer["目标周期节奏占比"])
    st.markdown(
        f"""
        - **周期识别**：根据生产库正式累计任务占年度任务的中位数，当前任务识别到
          **{meta["task_through_month"]}月**（观测占比 {meta["observed_task_ratio"]:.1%}，
          匹配标准累计节奏 {meta["matched_cumulative_ratio"]:.1%}）；默认规划下一期。
        - **任务进度**：生产库官方累计达成 ÷ {target_month}月底累计任务；
          本次目标周期新增节奏占年度任务 **{period_ratio:.1%}**。
        - **官方达成**：直接读取生产表 `signed_customer_monthly`
          的“累计业绩达成（计任务）”，不使用SI明细金额替代。
        - **SI型号数量**：仅统计“行业一级＝纯分销”且“外部型号≠discount”的实发数量；
          discount为返利扣减行，不进入客户备货结构。
        - **智能配单**：以客户近期SO、上述SI型号数量结构和SI−SO净流入代理指标分配金额；
          该代理指标不是库存。
        - **价格优先级**：最新报价单精确匹配 → 规范型号匹配 → 产品流向最新价 → SI期间加权单价。
        - **数据水位**：SI到 **{meta["si_latest_month"]}**（导入 {meta["si_imported_at"]}）；
          SO到 **{meta["so_asof"]}**；报价单日期 **{meta["quotation_date"]}**。
        - **安全边界**：生产业务库只读；页面编辑不落库。实际下单前仍需核对库存、在途、项目单与停产状态。
        """
    )
