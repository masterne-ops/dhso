#!/usr/bin/env python3
"""呆滞库存清理专项 · 周报 docx 生成

标准化数据(系统生成) + 用户填写的综述文字 → 合成 docx。
"""
from io import BytesIO
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0B, 0x25, 0x45)
ORANGE = RGBColor(0xC2, 0x41, 0x0C)
GREY = RGBColor(0x5A, 0x6B, 0x7B)
CN = '微软雅黑'


def _cn(run, font=CN):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexc)
    tcPr.append(shd)


def _bar(ratio, width=16):
    ratio = max(0.0, min(1.0, ratio))
    f = round(ratio * width)
    return '█' * f + '░' * (width - f) + f'  {ratio*100:.0f}%'


def _table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    hc = t.rows[0].cells
    for i, htext in enumerate(headers):
        _shade(hc[i], '0B2545')
        p = hc[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        r = p.add_run(htext)
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _cn(r)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            if i == 0:
                r.font.color.rgb = NAVY
            _cn(r)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t


def build_stale_report_docx(df, period, user_note=''):
    """df: stale_inventory; period: 本期清理月份(如 '2026-05'); user_note: 用户填的综述"""
    df = df.copy()
    df['_done'] = df['是否上线'] == 'Y'
    df['下单时单价'] = pd.to_numeric(df['下单时单价'], errors='coerce').fillna(0)
    df['_uamt'] = df['下单时单价'].where(~df['_done'], 0)
    total = len(df)
    done = int(df['_done'].sum())
    undone = total - done
    amt = df['下单时单价'].sum() / 1e4
    amt_done = df.loc[df['_done'], '下单时单价'].sum() / 1e4
    amt_undone = amt - amt_done

    doc = Document()
    ns = doc.styles['Normal']
    ns.font.name = CN
    ns.font.size = Pt(10.5)
    ns._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), CN)
    for sec in doc.sections:
        sec.top_margin = Inches(0.7)
        sec.bottom_margin = Inches(0.7)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    def para(text='', size=10.5, bold=False, color=None, align=None, after=6, before=0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        if text:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            if color:
                r.font.color.rgb = color
            _cn(r)
        return p

    def h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = NAVY
        _cn(r)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single')
        bot.set(qn('w:sz'), '6')
        bot.set(qn('w:space'), '4')
        bot.set(qn('w:color'), 'C2410C')
        pbdr.append(bot)
        pPr.append(pbdr)

    # ── 标题 ──
    para('呆滞库存清理专项 · 周报', 18, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=2, before=4)
    para(f'本期清理月份：{period}　|　代理商仓库 25 年初盘呆滞库存', 9, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=8)

    n = 0
    # ── 一、综述(用户填) ──
    if user_note and user_note.strip():
        n += 1
        h(f'{"一二三四五六"[n-1]}、本期综述')
        for line in user_note.strip().split('\n'):
            if line.strip():
                para(line.strip(), 10.5, after=4)

    # ── 整体进度 ──
    n += 1
    h(f'{"一二三四五六"[n-1]}、整体进度')
    para(f'呆滞总量 {total:,} 台（货值 {amt:.0f} 万）；已清理 {done:,} 台（{amt_done:.0f} 万，{done/total*100:.0f}%），'
         f'未清理 {undone:,} 台（{amt_undone:.0f} 万）。')
    p = para('整体清理　', 10.5, after=8)
    rb = p.add_run(_bar(done / total if total else 0, 24))
    rb.font.size = Pt(11)
    rb.font.color.rgb = ORANGE
    _cn(rb)

    # ── 本期清理情况 ──
    n += 1
    h(f'{"一二三四五六"[n-1]}、本期清理情况')
    cur = df[df['上线年月'] == period]
    para(f'{period} 共清理 {len(cur):,} 台 / {cur["下单时单价"].sum()/1e4:.1f} 万。各代理商本期清理：')
    cg = cur.groupby('代理商').agg(台数=('序列号', 'count'),
                                  货值=('下单时单价', lambda s: round(s.sum() / 1e4, 1))).reset_index().sort_values('台数', ascending=False)
    _table(doc, ['代理商', '本期清理台数', '货值(万)'],
           [[r['代理商'], r['台数'], r['货值']] for _, r in cg.iterrows()], widths=[3.4, 1.6, 1.4])
    para('本期清理的主要产品型号：', 10.5, True, NAVY, after=3, before=8)
    cm = cur.groupby('内部型号').size().reset_index(name='台数').sort_values('台数', ascending=False).head(12)
    _table(doc, ['内部型号', '清理台数'],
           [[r['内部型号'], r['台数']] for _, r in cm.iterrows()], widths=[5.0, 1.4])

    # ── 各代理商累计进度(进度条) ──
    n += 1
    h(f'{"一二三四五六"[n-1]}、各代理商累计清理进度')
    g = df.groupby('代理商').agg(呆滞=('序列号', 'count'), 已清=('_done', 'sum'),
                                剩余货值=('_uamt', lambda s: round(s.sum() / 1e4, 1))).reset_index()
    g['已清'] = g['已清'].astype(int)
    g['率'] = g['已清'] / g['呆滞']
    g = g.sort_values('呆滞', ascending=False)
    _table(doc, ['代理商', '呆滞', '已清', '清理进度', '剩余货值(万)'],
           [[r['代理商'], r['呆滞'], r['已清'], _bar(r['率'], 14), r['剩余货值']] for _, r in g.iterrows()],
           widths=[2.3, 0.7, 0.7, 2.4, 1.1])

    # ── 攻坚重点 ──
    n += 1
    h(f'{"一二三四五六"[n-1]}、攻坚重点（未清理）')
    und = df[~df['_done']]
    tm = und.groupby('内部型号').agg(剩余=('序列号', 'count'),
                                    货值=('下单时单价', lambda s: round(s.sum() / 1e4, 1)),
                                    主年份=('出库年份', lambda s: int(s.mode().iloc[0]) if not s.mode().dropna().empty else '—')).reset_index().sort_values('剩余', ascending=False).head(10)
    para('未清理最多的型号（换镜头/换POE改造 或 甩卖的攻坚对象）：', 10.5, after=3)
    _table(doc, ['内部型号', '剩余台数', '主出库年份', '剩余货值(万)'],
           [[r['内部型号'], r['剩余'], r['主年份'], r['货值']] for _, r in tm.iterrows()], widths=[3.4, 1.1, 1.3, 1.4])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
