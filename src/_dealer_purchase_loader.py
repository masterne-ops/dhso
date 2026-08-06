#!/usr/bin/env python3
"""代理商进货明细(工作簿3类导出)导入 — data-import 页「进货明细」区 + CLI 共用逻辑。

源文件: 单 Sheet 14 列,无年份/区县/计销售列,日期为「N月D日」文本:
  日期 / 业绩归属二级办 / 业务员姓名 / 下单客户 / 行业一级 / 产品一~四级 /
  内部型号 / 外部型号 / 求和项:实销（万) / 求和项:实发数量 / 求和项:计奖励

策略: **按年整年替换** dealer_purchase —— DELETE 数据年份=year + INSERT 文件全量,
其他年份不动。下单客户类型派生同 load_dealer_purchase.py 口径。
CLI 版: migrations/load_dealer_purchase_2026.py(逻辑一致,改动请两处同步)。
"""
import sqlite3
from datetime import datetime

import openpyxl

# 与 load_dealer_purchase.py / load_dealer_purchase_2026.py 保持一致
SMB_SUPPLEMENT = [
    '浙江名硕建设工程有限公司',
    '浙江歌呈智能科技有限公司',
    '微空间网络科技（湖州）有限公司',
    '湖州微空间网络科技有限公司',
]

REQUIRED_COLS = [
    '日期', '业绩归属二级办', '业务员姓名', '下单客户', '行业一级',
    '产品一级', '产品二级', '产品三级', '产品四级', '内部型号', '外部型号',
    '求和项:实销（万)', '求和项:实发数量', '求和项:计奖励',
]

_DB_COLS = ['数据年份', '月份', '日期原文', '业绩归属二级办', '业务员姓名', '下单客户', '下单客户类型',
            '下单客户区县', '行业一级', '产品一级', '产品二级', '产品三级', '产品四级',
            '内部型号', '外部型号', '实销万', '实发数量', '计返利', '计销售', '计奖励', '导入时间']


def _month_of(d):
    s = str(d or '')
    if '月' in s:
        try:
            return int(s.split('月')[0])
        except ValueError:
            return None
    return None


def _num(v):
    if v is None or v == '':
        return None
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return None


def parse_file(path, year):
    """解析文件 → (batch 行列表, 统计 dict)。缺必备列抛 ValueError。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else None for c in rows[0]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    miss = [e for e in REQUIRED_COLS if e not in idx]
    if miss:
        raise ValueError(f'文件缺必备列: {miss}(需要「工作簿3」类进货明细导出)')

    def s(v):
        t = str(v or '').strip()
        return t or None

    batch, by_month, amt = [], {}, 0.0
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    for r in rows[1:]:
        if not r or not r[idx['下单客户']]:
            continue
        m = _month_of(r[idx['日期']])
        by_month[m] = by_month.get(m, 0) + 1
        v = _num(r[idx['求和项:实销（万)']])
        amt += v or 0
        batch.append((
            year, m, s(r[idx['日期']]),
            s(r[idx['业绩归属二级办']]), s(r[idx['业务员姓名']]), s(r[idx['下单客户']]),
            None, None,
            s(r[idx['行业一级']]),
            s(r[idx['产品一级']]), s(r[idx['产品二级']]), s(r[idx['产品三级']]), s(r[idx['产品四级']]),
            s(r[idx['内部型号']]), s(r[idx['外部型号']]),
            v, _num(r[idx['求和项:实发数量']]),
            None, None, _num(r[idx['求和项:计奖励']]),
            now,
        ))
    stats = {'rows': len(batch), 'months': dict(sorted(by_month.items(), key=lambda x: (x[0] is None, x[0]))),
             'amt_wan': round(amt, 1)}
    return batch, stats


def apply_import(conn, batch, year):
    """单事务: DELETE 该年 + INSERT + 派生下单客户类型。返回统计 dict。"""
    old = conn.execute('SELECT COUNT(*) FROM dealer_purchase WHERE 数据年份=?', (year,)).fetchone()[0]
    ph = ','.join('?' * len(_DB_COLS))
    conn.execute('BEGIN')
    try:
        conn.execute('DELETE FROM dealer_purchase WHERE 数据年份=?', (year,))
        conn.executemany(f"INSERT INTO dealer_purchase ({','.join(_DB_COLS)}) VALUES ({ph})", batch)
        conn.execute("""
            UPDATE dealer_purchase SET 下单客户类型='一级代理商'
             WHERE 数据年份=? AND 下单客户 IN (SELECT DISTINCT "客户名称（无@办事处）" FROM dealer_si_snapshot)
        """, (year,))
        ph2 = ','.join('?' * len(SMB_SUPPLEMENT))
        conn.execute(f"UPDATE dealer_purchase SET 下单客户类型='一级代理商' "
                     f"WHERE 数据年份=? AND 下单客户 IN ({ph2})", (year, *SMB_SUPPLEMENT))
        conn.execute("""
            UPDATE dealer_purchase SET 下单客户类型='服务商'
             WHERE 数据年份=? AND 下单客户类型 IS NULL
               AND 下单客户 IN (SELECT DISTINCT 客户名称 FROM provider_contract)
        """, (year,))
        conn.execute("UPDATE dealer_purchase SET 下单客户类型='其他' "
                     "WHERE 数据年份=? AND 下单客户类型 IS NULL", (year,))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    types = dict(conn.execute(
        "SELECT 下单客户类型, COUNT(*) FROM dealer_purchase WHERE 数据年份=? GROUP BY 1", (year,)).fetchall())
    return {'deleted': old, 'inserted': len(batch), 'types': types}
